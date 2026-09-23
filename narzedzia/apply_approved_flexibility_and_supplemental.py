from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Cm


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)


def find_unique(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Oczekiwano jednego akapitu, znaleziono {len(matches)}: {text!r}")
    return matches[0]


def insert_after(paragraph, text: str, style=None):
    new_paragraph = paragraph._parent.add_paragraph(style=style or paragraph.style)
    new_paragraph.add_run(text)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


def replace_with_term(paragraph, number: str, term: str, definition: str):
    paragraph.clear()
    paragraph.add_run(number)
    paragraph.add_run(term).bold = True
    paragraph.add_run(definition)


def insert_definition_after(paragraph, number: str, term: str, definition: str):
    new_paragraph = paragraph._parent.add_paragraph(style=paragraph.style)
    new_paragraph.add_run(number)
    new_paragraph.add_run(term).bold = True
    new_paragraph.add_run(definition)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


document = Document(TARGET)

section_2_anchor = find_unique(
    document,
    "1. Celem zasad określonych w Regulaminie jest wyłonienie możliwie najsilniejszej "
    "reprezentacji Polski weteranów w szermierce, z zachowaniem przejrzystości procesu, "
    "równego traktowania zawodników oraz możliwości zweryfikowania podstaw podjętych decyzji.",
)
insert_after(
    section_2_anchor,
    "2. Regulamin zapewnia zawodnikom swobodę wyboru zawodów uwzględnianych w rankingu "
    "i nie wprowadza obowiązku udziału w zawodach krajowych, jednocześnie wykluczając, "
    "z wyjątkiem powołania uzupełniającego, możliwość powołania do reprezentacji zawodnika "
    "bez zweryfikowanego dorobku punktowego w rankingu indywidualnym.",
)

definition_anchor = find_unique(
    document,
    "3. wygaszanie kalendarzowe – sposób uwzględniania wyników zawodów organizowanych przez "
    "EVF, zgodnie z którym uzyskane punkty pozostają w rankingu przez okres właściwy dla danego "
    "rodzaju zawodów, liczony od dnia zakończenia zawodów, niezależnie od tego, czy zawody mają "
    "odpowiednik w kolejnym sezonie, a po upływie tego okresu przestają być uwzględniane.",
)
replace_with_term(
    definition_anchor,
    "3. ",
    "wygaszanie kalendarzowe",
    " – sposób uwzględniania wyników zawodów organizowanych przez EVF, zgodnie z którym "
    "uzyskane punkty pozostają w rankingu przez okres właściwy dla danego rodzaju zawodów, "
    "liczony od dnia zakończenia zawodów, niezależnie od tego, czy zawody mają odpowiednik "
    "w kolejnym sezonie, a po upływie tego okresu przestają być uwzględniane;",
)
definition_4 = insert_definition_after(
    definition_anchor,
    "4. ",
    "limitowane zawody międzynarodowe",
    " – indywidualne lub drużynowe zawody międzynarodowe rangi mistrzowskiej, w których liczba "
    "zawodników reprezentujących dane państwo jest ograniczona przepisami organizatora zawodów;",
)
insert_definition_after(
    definition_4,
    "5. ",
    "powołanie uzupełniające",
    " – wyjątkowe powołanie do startu drużynowego zawodnika nieposiadającego dodatniego dorobku "
    "punktowego, dokonywane wyłącznie w celu uzupełnienia braków uniemożliwiających wystawienie "
    "kompletnej i zgodnej z przepisami drużyny.",
)

section_4_anchor = find_unique(
    document,
    "2. Miejsce w rankingu indywidualnym nie stanowi samoistnej podstawy powołania do składu "
    "drużyny. Skład drużyny ustala się zgodnie z zasadami określonymi w rozdziale IV.",
)
section_4_3 = insert_after(
    section_4_anchor,
    "3. Zawodnik samodzielnie decyduje, w których zawodach wymienionych w § 5 bierze udział. "
    "Udział w określonych zawodach lub cyklu zawodów nie stanowi warunku powołania do reprezentacji.",
)
section_4_4 = insert_after(
    section_4_3,
    "4. Z zastrzeżeniem ust. 5, do reprezentacji Polski na limitowane zawody międzynarodowe może "
    "zostać powołany wyłącznie zawodnik, który wystartował w co najmniej jednych zawodach "
    "uwzględnianych w rankingu indywidualnym i uzyskał w nim dorobek większy niż zero punktów.",
)
section_4_5 = insert_after(
    section_4_4,
    "5. W przypadku gdy nie jest możliwe skompletowanie drużyny z zawodników posiadających dodatni "
    "dorobek punktowy, dopuszcza się powołanie uzupełniające zawodnika bez punktów rankingowych, "
    "jeżeli łącznie:",
)

conditions = [
    "1) powołanie jest niezbędne do wystawienia kompletnej drużyny zgodnej z przepisami zawodów;",
    "2) możliwość startu zaoferowano wszystkim dostępnym zawodnikom posiadającym dodatni dorobek "
    "punktowy, spełniającym wymagania dotyczące broni, płci i kategorii wiekowej;",
    "3) żaden z zawodników, o których mowa w pkt 2, nie potwierdził gotowości do udziału;",
    "4) powoływana osoba spełnia wszystkie formalne warunki udziału określone przez organizatora zawodów;",
    "5) powołanie nie pozbawia miejsca zawodnika posiadającego dodatni dorobek punktowy i gotowego do udziału;",
    "6) przyczyny powołania zostaną wskazane w pisemnym uzasadnieniu nominacji.",
]
cursor = section_4_5
for condition in conditions:
    cursor = insert_after(cursor, condition)
    cursor.paragraph_format.left_indent = Cm(0.65)
    cursor.paragraph_format.first_line_indent = Cm(-0.4)

section_5_anchor = find_unique(
    document,
    "8) opublikowane wyniki zawierają dane niezbędne do ich zweryfikowania oraz obliczenia punktów rankingowych.",
)
insert_after(
    section_5_anchor,
    "3. Spośród wyników uzyskanych w zawodach, o których mowa w ust. 1 pkt 1, w rankingu "
    "uwzględnia się trzy wyniki, za które zawodnik uzyskał najwyższą liczbę punktów rankingowych.",
)

document.core_properties.comments = (
    "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1–§ 5; "
    "pozostała treść normatywna nie została jeszcze uzgodniona."
)

timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-zasada-elastycznosci-{timestamp}.docx")
shutil.copy2(TARGET, backup)
temporary = TARGET.with_suffix(".tmp.docx")
document.save(temporary)
os.replace(temporary, TARGET)

print(f"Zapisano: {TARGET}")
print(f"Kopia: {backup}")
