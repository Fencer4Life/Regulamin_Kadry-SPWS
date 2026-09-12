from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-reprezentacji-Polski-weteranów-w-szermierce_2026.docx"
)

DEFINITIONS = (
    (
        "1. ",
        "ranking hybrydowy",
        " – ranking indywidualny, w którym wyniki są aktualizowane z zastosowaniem "
        "aktualizacji krokowej albo wygaszania kalendarzowego, zależnie od organizatora "
        "i rodzaju zawodów;",
    ),
    (
        "2. ",
        "aktualizacja krokowa",
        " – sposób aktualizacji wyników zawodów organizowanych przez SPWS, PZSz lub FIE, "
        "zgodnie z którym punkty za zawody z poprzedniego sezonu pozostają w rankingu do "
        "czasu uwzględnienia punktów za odpowiadające im zawody w sezonie bieżącym, po czym "
        "poprzedni wynik przestaje być uwzględniany;",
    ),
    (
        "3. ",
        "wygaszanie kalendarzowe",
        " – sposób uwzględniania wyników zawodów organizowanych przez EVF, zgodnie z którym "
        "uzyskane punkty pozostają w rankingu przez okres właściwy dla danego rodzaju "
        "zawodów, liczony od dnia zakończenia zawodów, niezależnie od tego, czy zawody mają "
        "odpowiednik w kolejnym sezonie, a po upływie tego okresu przestają być uwzględniane.",
    ),
)


def set_definition(paragraph: Paragraph, number: str, term: str, explanation: str) -> None:
    paragraph.clear()
    paragraph.style = "Normal"
    paragraph.add_run(number)
    paragraph.add_run(term).bold = True
    paragraph.add_run(explanation)


def insert_after(paragraph: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    return Paragraph(element, paragraph._parent)


def main() -> None:
    if not TARGET.is_file():
        raise SystemExit(f"Nie znaleziono dokumentu docelowego: {TARGET}")

    document = Document(TARGET)
    paragraphs = document.paragraphs
    expected = {
        29: "§ 3. Definicje",
        30: "1. [Definicje pojęć używanych w Regulaminie.]",
        31: "ROZDZIAŁ II",
    }
    for index, text in expected.items():
        if paragraphs[index].text != text:
            raise SystemExit(
                f"Dokument zmienił się w nieoczekiwanym miejscu {index}: "
                f"{paragraphs[index].text!r} zamiast {text!r}. Nie zapisano zmian."
            )

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-paragrafem-3-{timestamp}.docx")
    shutil.copy2(TARGET, backup)

    current = paragraphs[30]
    set_definition(current, *DEFINITIONS[0])
    for definition in DEFINITIONS[1:]:
        current = insert_after(current)
        set_definition(current, *definition)

    document.core_properties.comments = (
        "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1–§ 4; "
        "pozostała treść normatywna nie została jeszcze uzgodniona."
    )

    temporary = TARGET.with_name(f".{TARGET.name}.tmp")
    document.save(temporary)
    os.replace(temporary, TARGET)

    print(f"TARGET={TARGET}")
    print(f"BACKUP={backup}")


if __name__ == "__main__":
    main()
