from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)

PARAGRAPH_1 = (
    "1. Ranking indywidualny stanowi podstawę kwalifikowania zawodników do startów "
    "indywidualnych oraz wyłonienia puli kandydatów do składu drużyny."
)

PARAGRAPH_2 = (
    "2. Miejsce w rankingu indywidualnym nie stanowi samoistnej podstawy powołania do "
    "składu drużyny. Skład drużyny ustala się zgodnie z zasadami określonymi w rozdziale IV."
)


def insert_after(paragraph: Paragraph, text: str, style: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def main() -> None:
    if not TARGET.is_file():
        raise SystemExit(f"Nie znaleziono dokumentu docelowego: {TARGET}")

    document = Document(TARGET)
    paragraphs = document.paragraphs
    expected = {
        31: "ROZDZIAŁ II",
        32: "Ranking indywidualny",
        33: "§ 4",
        34: "1. [Cel i znaczenie rankingu indywidualnego.]",
        35: "§ 5",
    }
    for index, text in expected.items():
        if paragraphs[index].text != text:
            raise SystemExit(
                f"Dokument zmienił się w nieoczekiwanym miejscu {index}: "
                f"{paragraphs[index].text!r} zamiast {text!r}. Nie zapisano zmian."
            )

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-paragrafem-4-{timestamp}.docx")
    shutil.copy2(TARGET, backup)

    paragraphs[33].text = "§ 4. Rola rankingu indywidualnego"
    paragraphs[33].style = "Paragraf"
    paragraphs[34].text = PARAGRAPH_1
    paragraphs[34].style = "Normal"
    insert_after(paragraphs[34], PARAGRAPH_2, "Normal")

    document.core_properties.comments = (
        "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1, § 2 oraz § 4; "
        "pozostała treść normatywna nie została jeszcze uzgodniona."
    )

    temporary = TARGET.with_name(f".{TARGET.name}.tmp")
    document.save(temporary)
    os.replace(temporary, TARGET)

    print(f"TARGET={TARGET}")
    print(f"BACKUP={backup}")


if __name__ == "__main__":
    main()
