from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)

CONTENT = (
    "1. W sezonie 2026/2027 w rankingu indywidualnym uwzględnia się wyniki "
    "uzyskane w następujących cyklach i zawodach:",
    "1) Pucharze Polski Weteranów w Szermierce oraz Mistrzostwach Polski Weteranów "
    "w Szermierce, organizowanych przez SPWS albo przez inny podmiot na zasadach "
    "określonych w ust. 2;",
    "2) Pucharze Polski Seniorów w Szermierce oraz Mistrzostwach Polski Seniorów "
    "w Szermierce, ujętych w oficjalnym kalendarzu Polskiego Związku Szermierczego;",
    "3) Pucharze Europy Weteranów w Szermierce oraz Mistrzostwach Europy Weteranów "
    "w Szermierce, ujętych w oficjalnym kalendarzu European Veterans Fencing;",
    "4) Pucharze Świata Weteranów w Szermierce – o ile zawody tego cyklu zostaną "
    "rozegrane – oraz Mistrzostwach Świata Weteranów w Szermierce, ujętych w "
    "oficjalnym kalendarzu Międzynarodowej Federacji Szermierczej.",
    "2. Zawody organizowane przez podmiot inny niż SPWS mogą zostać zaliczone do "
    "cyklu Pucharu Polski Weteranów w Szermierce, jeżeli łącznie spełniają następujące "
    "warunki:",
    "1) przed ich rozegraniem zostały zatwierdzone przez SPWS i ujęte w opublikowanym "
    "kalendarzu cyklu;",
    "2) komunikat organizacyjny zawodów został opublikowany w ogólnodostępnym serwisie "
    "internetowym, bez konieczności logowania, najpóźniej 30 dni przed pierwszym dniem "
    "zawodów;",
    "3) prawo zgłoszenia do udziału przysługuje wszystkim zawodnikom spełniającym "
    "kryterium wieku właściwe dla weteranów szermierki, niezależnie od obywatelstwa, "
    "miejsca zamieszkania, przynależności klubowej oraz państwa, w którym uprawiają "
    "szermierkę, pod warunkiem spełnienia jednakowych dla wszystkich uczestników "
    "wymagań określonych w komunikacie organizacyjnym;",
    "4) zawody są rozgrywane według formuły sportowej zgodnej z regulaminami lub "
    "przepisami zawodów stosowanymi przez SPWS, PZSz, EVF albo FIE, przy czym właściwa "
    "formuła musi zostać wskazana w komunikacie organizacyjnym;",
    "5) zawody są rozgrywane w kategoriach wieku, broni i płci określonych w komunikacie "
    "organizacyjnym i zgodnych z wybraną formułą sportową;",
    "6) po zakończeniu zawodów organizator publikuje w serwisie internetowym kompletne "
    "protokoły wszystkich rozegranych walk, obejmujące fazę grupową oraz eliminację "
    "bezpośrednią, wraz z klasyfikacją końcową i liczbą uczestników, zapewniając publiczny "
    "i nieodpłatny dostęp do tych materiałów;",
    "7) protokoły i wyniki, o których mowa w pkt 6, pozostają dostępne online przez co "
    "najmniej 36 miesięcy od dnia zakończenia zawodów;",
    "8) opublikowane wyniki zawierają dane niezbędne do ich zweryfikowania oraz obliczenia "
    "punktów rankingowych.",
)


def insert_after(paragraph: Paragraph, text: str, nested: bool = False) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = "Normal"
    inserted.add_run(text)
    if nested:
        inserted.paragraph_format.left_indent = Cm(0.65)
        inserted.paragraph_format.first_line_indent = Cm(-0.4)
    return inserted


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def main() -> None:
    if not TARGET.is_file():
        raise SystemExit(f"Nie znaleziono dokumentu docelowego: {TARGET}")

    document = Document(TARGET)
    paragraphs = document.paragraphs
    expected = {
        38: "§ 5",
        39: "1. [Kategorie, zawody i wyniki uwzględniane w rankingu.]",
        40: "2. [Zasady punktacji i publikacji rankingu.]",
        42: "ROZDZIAŁ III",
    }
    for index, text in expected.items():
        if paragraphs[index].text != text:
            raise SystemExit(
                f"Dokument zmienił się w nieoczekiwanym miejscu {index}: "
                f"{paragraphs[index].text!r} zamiast {text!r}. Nie zapisano zmian."
            )

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-paragrafem-5-{timestamp}.docx")
    shutil.copy2(TARGET, backup)

    paragraphs[38].text = "§ 5. Zawody uwzględniane w rankingu"
    paragraphs[38].style = "Paragraf"

    first = paragraphs[39]
    first.text = CONTENT[0]
    first.style = "Normal"
    current = first
    for index, text in enumerate(CONTENT[1:], start=1):
        # Pozycje 1–4 należą do ust. 1, a pozycje 6–13 do ust. 2.
        nested = 1 <= index <= 4 or index >= 6
        current = insert_after(current, text, nested=nested)

    remove_paragraph(paragraphs[40])

    document.core_properties.comments = (
        "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1–§ 5; "
        "pozostała treść normatywna nie została jeszcze uzgodniona."
    )

    temporary = TARGET.with_name(f".{TARGET.name}.tmp")
    document.save(temporary)
    os.replace(temporary, TARGET)

    print(f"TARGET={TARGET}")
    print(f"BACKUP={backup}")


if __name__ == "__main__":
    main()
