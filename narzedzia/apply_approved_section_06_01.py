from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)


def find_unique(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Oczekiwano jednego akapitu, znaleziono {len(matches)}: {text!r}")
    return matches[0]


def insert_after(paragraph, text: str, style=None):
    new_paragraph = paragraph._parent.add_paragraph(style=style or paragraph.style)
    new_paragraph.add_run(text)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


document = Document(TARGET)

anchor = find_unique(
    document,
    "3. Spośród wyników uzyskanych w zawodach, o których mowa w ust. 1 pkt 1, w rankingu "
    "uwzględnia się trzy wyniki, za które zawodnik uzyskał najwyższą liczbę punktów rankingowych.",
)

section_title = insert_after(anchor, "§ 6. Zasady obliczania punktów", style="Paragraf")
insert_after(
    section_title,
    "1. Liczba punktów przyznawanych za wynik uzyskany w zawodach zależy od zajętego miejsca, "
    "liczby zawodników w stawce oraz współczynnika rangi określonego dla danego rodzaju zawodów.",
    style="Normal",
)

for old_number in range(11, 5, -1):
    paragraph = find_unique(document, f"§ {old_number}")
    paragraph.text = f"§ {old_number + 1}"

document.core_properties.comments = (
    "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1–§ 6; "
    "pozostała treść normatywna nie została jeszcze uzgodniona."
)

timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-paragrafem-06-01-{timestamp}.docx")
shutil.copy2(TARGET, backup)
temporary = TARGET.with_suffix(".tmp.docx")
document.save(temporary)
os.replace(temporary, TARGET)

print(f"Zapisano: {TARGET}")
print(f"Kopia: {backup}")
