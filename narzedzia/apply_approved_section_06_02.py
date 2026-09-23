from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)

ROWS = [
    ("PPW", "Puchar Polski Weteranów w szermierce", "1,0"),
    ("MPW", "Mistrzostwa Polski Weteranów w szermierce", "1,2"),
    ("PPS", "Puchar Polski Seniorów w szermierce", "1,1"),
    ("MPS", "Mistrzostwa Polski Seniorów w szermierce", "1,3"),
    ("PEW", "Puchar Europy Weteranów w szermierce", "1,1"),
    ("MEW", "Mistrzostwa Europy Weteranów w szermierce", "1,3"),
    ("PSW", "Puchar Świata Weteranów w szermierce", "1,1"),
    ("MŚW", "Mistrzostwa Świata Weteranów w szermierce", "1,4"),
]


def find_unique(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Oczekiwano jednego akapitu, znaleziono {len(matches)}: {text!r}")
    return matches[0]


def insert_after(paragraph, text: str, style=None):
    new_paragraph = paragraph._parent.add_paragraph(style=style or paragraph.style)
    new_paragraph.add_run(text)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


def set_cell_shading(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


document = Document(TARGET)
anchor = find_unique(
    document,
    "1. Liczba punktów przyznawanych za wynik uzyskany w zawodach zależy od zajętego miejsca, "
    "liczby zawodników w stawce oraz współczynnika rangi określonego dla danego rodzaju zawodów.",
)

intro = insert_after(
    anchor,
    "2. W sezonie 2026/2027 stosuje się następujące współczynniki rangi zawodów:",
    style="Normal",
)

table = document.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.autofit = False
table.columns[0].width = Cm(2.2)
table.columns[1].width = Cm(11.2)
table.columns[2].width = Cm(3.2)
table._tbl.getparent().remove(table._tbl)
intro._p.addnext(table._tbl)

headers = ("Skrót", "Rodzaj zawodów", "Współczynnik")
header_row = table.rows[0]
set_repeat_table_header(header_row)
for index, (cell, label) in enumerate(zip(header_row.cells, headers)):
    cell.width = (Cm(2.2), Cm(11.2), Cm(3.2))[index]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "173F73")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index != 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(label)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    run.font.size = Pt(9)

for row_index, values in enumerate(ROWS, start=1):
    cells = table.add_row().cells
    for column_index, (cell, value) in enumerate(zip(cells, values)):
        cell.width = (Cm(2.2), Cm(11.2), Cm(3.2))[column_index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index % 2 == 0:
            set_cell_shading(cell, "EAF2FB")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index != 1 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        run.font.size = Pt(9)
        if column_index in (0, 2):
            run.bold = True

document.core_properties.comments = (
    "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1–§ 6; "
    "pozostała treść normatywna nie została jeszcze uzgodniona."
)

timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-paragrafem-06-02-{timestamp}.docx")
shutil.copy2(TARGET, backup)
temporary = TARGET.with_suffix(".tmp.docx")
document.save(temporary)
os.replace(temporary, TARGET)

print(f"Zapisano: {TARGET}")
print(f"Kopia: {backup}")
