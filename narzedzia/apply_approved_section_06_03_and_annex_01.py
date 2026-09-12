from __future__ import annotations

import math
import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-reprezentacji-Polski-weteranów-w-szermierce_2026.docx"
)

PLACE_BLOCKS = ((1, 10), (11, 20), (21, 30), (31, 40))
MIN_PARTICIPANTS = 4
MAX_PARTICIPANTS = 40
RANK_COEFFICIENT = 1.0


def find_unique(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Oczekiwano jednego akapitu, znaleziono {len(matches)}: {text!r}")
    return matches[0]


def is_power_of_two(value: int) -> bool:
    return value > 0 and (value & (value - 1)) == 0


def score(place: int, participants: int) -> float:
    if place < 1 or place > participants:
        raise ValueError("Miejsce musi należeć do stawki zawodników")
    base = min(50.0, 10.0 * math.log2(max(2, participants)))
    place_points = base - (base - 1.0) * math.log(place) / math.log(participants)
    won_rounds = max(
        0,
        math.floor(math.log2(participants))
        - math.ceil(math.log2(place))
        + (0 if is_power_of_two(participants) else 1),
    )
    podium_factor = {1: 3, 2: 2, 3: 1}.get(place, 0)
    podium = podium_factor * 3.0 * participants ** (1.0 / 3.0)
    return (place_points + won_rounds * 10.0 + podium) * RANK_COEFFICIENT


def format_points(value: float) -> str:
    return f"{value:.1f}".replace(".", ",")


def shade(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def prevent_row_split(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def format_cell(cell, text: str, *, header=False, first_column=False, alternate=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        shade(cell, "173F73")
    elif first_column:
        shade(cell, "EAF2FB")
    elif alternate:
        shade(cell, "F6F8FB")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.size = Pt(7.5)
    run.bold = header or first_column
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)


document = Document(TARGET)

find_unique(
    document,
    "2. W sezonie 2026/2027 stosuje się następujące współczynniki rangi zawodów:",
)
coefficient_tables = [
    table
    for table in document.tables
    if table.cell(0, 0).text == "Skrót" and table.cell(0, 1).text == "Rodzaj zawodów"
]
if len(coefficient_tables) != 1:
    raise RuntimeError(f"Oczekiwano jednej tabeli współczynników, znaleziono {len(coefficient_tables)}")
section_6_3 = document.add_paragraph(style="Normal")
section_6_3.add_run(
    "3. Tabela punktacji dla stawki liczącej od 4 do 40 zawodników stanowi załącznik nr 1 "
    "do Regulaminu. Pełną tabelę punktacji dla stawki liczącej od 4 do 300 zawodników SPWS "
    "publikuje na swojej stronie internetowej."
)
section_6_3._p.getparent().remove(section_6_3._p)
coefficient_tables[0]._tbl.addnext(section_6_3._p)

history_anchor = find_unique(document, "Historia wersji")
annex_elements = []

page_before = document.add_paragraph()
page_before.add_run().add_break(WD_BREAK.PAGE)
annex_elements.append(page_before._p)

label = document.add_paragraph(style="Etykieta rozdzialu")
label.alignment = WD_ALIGN_PARAGRAPH.CENTER
label.add_run("ZAŁĄCZNIK NR 1")
annex_elements.append(label._p)

title = document.add_paragraph(style="Heading 1")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Tabela punktacji Pucharu Polski Weteranów w szermierce")
annex_elements.append(title._p)

subtitle = document.add_paragraph(style="Normal")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run(
    "Punkty rankingowe za miejsce zdobyte w stawce liczącej od 4 do 40 zawodników "
    "— sezon 2026/2027"
).italic = True
annex_elements.append(subtitle._p)

coefficient = document.add_paragraph(style="Normal")
coefficient.alignment = WD_ALIGN_PARAGRAPH.CENTER
coefficient.add_run("Współczynnik rangi PPW: 1,0").bold = True
annex_elements.append(coefficient._p)

for first_place, last_place in PLACE_BLOCKS:
    block_heading = document.add_paragraph(style="Heading 2")
    block_heading.add_run(f"Miejsca {first_place}–{last_place}")
    annex_elements.append(block_heading._p)

    places = list(range(first_place, last_place + 1))
    first_participant = max(MIN_PARTICIPANTS, first_place)
    table = document.add_table(rows=1, cols=1 + len(places))
    table.style = "Table Grid"
    table.autofit = False
    first_width = Cm(3.0)
    score_width = Cm(1.25)

    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    header.cells[0].width = first_width
    format_cell(header.cells[0], "Liczba zawodników w stawce", header=True)
    for index, place in enumerate(places, start=1):
        header.cells[index].width = score_width
        format_cell(header.cells[index], str(place), header=True)

    for row_number, participants in enumerate(range(first_participant, MAX_PARTICIPANTS + 1)):
        row = table.add_row()
        prevent_row_split(row)
        row.cells[0].width = first_width
        format_cell(row.cells[0], str(participants), first_column=True)
        for column_index, place in enumerate(places, start=1):
            row.cells[column_index].width = score_width
            value = format_points(score(place, participants)) if place <= participants else "—"
            format_cell(row.cells[column_index], value, alternate=row_number % 2 == 1)
    annex_elements.append(table._tbl)

note = document.add_paragraph(style="Normal")
note.paragraph_format.space_before = Pt(8)
note.add_run(
    "Pełna tabela dla stawek liczących od 4 do 300 zawodników oraz kalkulator punktów "
    "są publikowane przez SPWS na stronie internetowej."
)
annex_elements.append(note._p)

page_after = document.add_paragraph()
page_after.add_run().add_break(WD_BREAK.PAGE)
annex_elements.append(page_after._p)

for element in annex_elements:
    element.getparent().remove(element)
    history_anchor._p.addprevious(element)

document.core_properties.comments = (
    "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1–§ 6 oraz załącznik nr 1; "
    "pozostała treść normatywna nie została jeszcze uzgodniona."
)

timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-paragrafem-06-03-i-zalacznikiem-01-{timestamp}.docx")
shutil.copy2(TARGET, backup)
temporary = TARGET.with_suffix(".tmp.docx")
document.save(temporary)
os.replace(temporary, TARGET)

print(f"Zapisano: {TARGET}")
print(f"Kopia: {backup}")
