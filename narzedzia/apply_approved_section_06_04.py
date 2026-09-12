from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-reprezentacji-Polski-weteranów-w-szermierce_2026.docx"
)


def find_unique(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Oczekiwano jednego akapitu, znaleziono {len(matches)}: {text!r}")
    return matches[0]


def insert_after(paragraph, text: str, style=None):
    new_paragraph = paragraph._parent.add_paragraph(style=style or paragraph.style)
    new_paragraph.add_run(text)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


document = Document(TARGET)
anchor = find_unique(
    document,
    "3. Tabela punktacji dla stawki liczącej od 4 do 40 zawodników stanowi załącznik nr 1 "
    "do Regulaminu. Pełną tabelę punktacji dla stawki liczącej od 4 do 300 zawodników SPWS "
    "publikuje na swojej stronie internetowej.",
)
insert_after(
    anchor,
    "4. SPWS zapewnia publiczny i nieodpłatny dostęp do kalkulatora punktów, umożliwiającego "
    "sprawdzenie liczby punktów przyznawanych za konkretny wynik uzyskany w zawodach.",
    style="Normal",
)

document.core_properties.comments = (
    "Projekt do konsultacji. Zatwierdzono i wprowadzono § 1–§ 6 oraz załącznik nr 1; "
    "pozostała treść normatywna nie została jeszcze uzgodniona."
)

timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-paragrafem-06-04-{timestamp}.docx")
shutil.copy2(TARGET, backup)
temporary = TARGET.with_suffix(".tmp.docx")
document.save(temporary)
os.replace(temporary, TARGET)

print(f"Zapisano: {TARGET}")
print(f"Kopia: {backup}")
