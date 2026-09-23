from __future__ import annotations

import math
import os
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TARGET = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)

PLACE_BLOCKS = ((1, 10), (11, 20), (21, 30), (31, 40))
PARTICIPANT_CHUNK_SIZE = 17
MIN_PARTICIPANTS = 4
MAX_PARTICIPANTS = 40


def find_unique(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Oczekiwano jednego akapitu, znaleziono {len(matches)}: {text!r}")
    return matches[0]


def is_power_of_two(value: int) -> bool:
    return value > 0 and (value & (value - 1)) == 0


def score(place: int, participants: int) -> float:
    base = min(50.0, 10.0 * math.log2(max(2, participants)))
    place_points = base - (base - 1.0) * math.log(place) / math.log(participants)
    won_rounds = max(
        0,
        math.floor(math.log2(participants))
        - math.ceil(math.log2(place))
        + (0 if is_power_of_two(participants) else 1),
    )
    podium_factor = {1: 3, 2: 2, 3: 1}.get(place, 0)
    podium = podium_factor * 3.0 * participants ** (1.0 / 3.0)
    return place_points + won_rounds * 10.0 + podium


def format_points(value: float) -> str:
    return f"{value:.1f}".replace(".", ",")


def shade(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is None:
        marker = OxmlElement("w:tblHeader")
        marker.set(qn("w:val"), "true")
        properties.append(marker)


def set_row_indivisible(row):
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def format_annex_cell(cell, text: str, *, header=False, first_column=False, alternate=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        shade(cell, "173F73")
    elif first_column:
        shade(cell, "EAF2FB")
    elif alternate:
        shade(cell, "F6F8FB")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.size = Pt(7.5)
    run.bold = header or first_column
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)


def append_before(document: Document, anchor, element):
    element.getparent().remove(element)
    anchor._p.addprevious(element)


def build_annex(document: Document, history_anchor):
    label = document.add_paragraph(style="Etykieta rozdzialu")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.add_run("ZAŁĄCZNIK NR 1")
    append_before(document, history_anchor, label._p)

    title = document.add_paragraph(style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Tabela punktacji Pucharu Polski Weteranów w szermierce")
    append_before(document, history_anchor, title._p)

    subtitle = document.add_paragraph(style="Normal")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Punkty rankingowe za miejsce zdobyte w stawce liczącej od 4 do 40 zawodników "
        "— sezon 2026/2027"
    ).italic = True
    append_before(document, history_anchor, subtitle._p)

    coefficient = document.add_paragraph(style="Normal")
    coefficient.alignment = WD_ALIGN_PARAGRAPH.CENTER
    coefficient.add_run("Współczynnik rangi PPW: 1,0").bold = True
    append_before(document, history_anchor, coefficient._p)

    for first_place, last_place in PLACE_BLOCKS:
        places = list(range(first_place, last_place + 1))
        first_participant = max(MIN_PARTICIPANTS, first_place)
        participants = list(range(first_participant, MAX_PARTICIPANTS + 1))

        for start in range(0, len(participants), PARTICIPANT_CHUNK_SIZE):
            chunk = participants[start : start + PARTICIPANT_CHUNK_SIZE]
            heading = document.add_paragraph(style="Heading 2")
            heading.paragraph_format.keep_with_next = True
            heading.add_run(
                f"Miejsca {first_place}–{last_place} · stawka {chunk[0]}–{chunk[-1]} zawodników"
            )
            append_before(document, history_anchor, heading._p)

            table = document.add_table(rows=1, cols=1 + len(places))
            table.style = "Table Grid"
            table.autofit = False
            first_width = Cm(3.0)
            score_width = Cm(1.25)

            header = table.rows[0]
            set_repeat_header(header)
            header.cells[0].width = first_width
            format_annex_cell(header.cells[0], "Liczba zawodników w stawce", header=True)
            for index, place in enumerate(places, start=1):
                header.cells[index].width = score_width
                format_annex_cell(header.cells[index], str(place), header=True)

            for row_number, participant_count in enumerate(chunk):
                row = table.add_row()
                row.cells[0].width = first_width
                format_annex_cell(row.cells[0], str(participant_count), first_column=True)
                for column_index, place in enumerate(places, start=1):
                    row.cells[column_index].width = score_width
                    value = format_points(score(place, participant_count)) if place <= participant_count else "—"
                    format_annex_cell(row.cells[column_index], value, alternate=row_number % 2 == 1)
            append_before(document, history_anchor, table._tbl)

    note = document.add_paragraph(style="Normal")
    note.paragraph_format.space_before = Pt(8)
    note.add_run(
        "Pełna tabela dla stawek liczących od 4 do 300 zawodników oraz kalkulator punktów "
        "są publikowane przez SPWS na stronie internetowej."
    )
    append_before(document, history_anchor, note._p)

    page_after = document.add_paragraph()
    page_after.add_run().add_break(WD_BREAK.PAGE)
    append_before(document, history_anchor, page_after._p)


document = Document(TARGET)

# Każdy rozdział rozpoczyna się na nowej stronie.
for paragraph in document.paragraphs:
    if paragraph.text.startswith("ROZDZIAŁ "):
        paragraph.paragraph_format.page_break_before = True

# Pierwszy paragraf pozostaje z tytułem rozdziału. Każdy następny paragraf
# w tym samym rozdziale rozpoczyna kolejną stronę, z wyjątkiem § 2.
section_seen_in_chapter = False
for paragraph in document.paragraphs:
    if paragraph.text.startswith("ROZDZIAŁ "):
        section_seen_in_chapter = False
    elif paragraph.text.startswith("§ "):
        if section_seen_in_chapter:
            paragraph.paragraph_format.page_break_before = True
        section_seen_in_chapter = True

# § 1 i § 2 stanowią jeden blok strony. Stały odstęp przed § 2 oddziela
# oba paragrafy wizualnie, a § 3 nadal rozpoczyna się na nowej stronie.
section_1 = find_unique(document, "§ 1. Przedmiot regulaminu")
section_2 = find_unique(document, "§ 2. Cel regulaminu")
section_3 = find_unique(document, "§ 3. Definicje")
section_2.paragraph_format.page_break_before = False
section_2.paragraph_format.space_before = Pt(24)

paragraphs = document.paragraphs
section_1_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph.text == section_1.text)
section_3_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph.text == section_3.text)
for paragraph in paragraphs[section_1_index : section_3_index - 1]:
    paragraph.paragraph_format.keep_with_next = True
paragraphs[section_3_index - 1].paragraph_format.keep_with_next = False

# Zastąp dotychczasowe duże tabele załącznika blokami mieszczącymi się na jednej stronie.
annex_start = find_unique(document, "ZAŁĄCZNIK NR 1")
history_anchor = find_unique(document, "Historia wersji")
body = document.element.body
children = list(body.iterchildren())
start_index = children.index(annex_start._p)
end_index = children.index(history_anchor._p)
for element in children[start_index:end_index]:
    body.remove(element)
build_annex(document, history_anchor)

# Word nie ma pojedynczego przełącznika „nie dziel tabeli”. Łączymy wszystkie
# jej wiersze keep-with-next i dodatkowo zabraniamy dzielenia pojedynczych wierszy.
for table in document.tables:
    for row_index, row in enumerate(table.rows):
        set_row_indivisible(row)
        keep_with_next = row_index < len(table.rows) - 1
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = keep_with_next

# Akapit bezpośrednio wprowadzający tabelę współczynników pozostaje z tabelą.
find_unique(
    document,
    "2. W sezonie 2026/2027 stosuje się następujące współczynniki rangi zawodów:",
).paragraph_format.keep_with_next = True

timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
backup = TARGET.with_name(f"{TARGET.stem}.backup-przed-poprawa-stronicowania-{timestamp}.docx")
shutil.copy2(TARGET, backup)
temporary = TARGET.with_suffix(".tmp.docx")
document.save(temporary)
os.replace(temporary, TARGET)

print(f"Zapisano: {TARGET}")
print(f"Kopia: {backup}")
