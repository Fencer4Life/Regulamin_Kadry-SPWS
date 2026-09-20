from __future__ import annotations

import argparse
import math
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

from narzedzia.docx_model import (
    RegulationDocument,
    TableBlock,
    ZtpUnit,
    parse_regulation_source,
)
from narzedzia.normalize_regulamin_markdown import resolve_references
from narzedzia.generate_regulamin_docx import (
    BLUE,
    LIGHT,
    LINE,
    MUTED,
    add_cover,
    add_field,
    add_header_footer,
    configure_page,
    configure_styles,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_table_borders,
)


PLACE_BLOCKS = ((1, 10), (11, 20), (21, 30), (31, 40))
PARTICIPANT_CHUNK_SIZE = 17
MIN_PARTICIPANTS = 4
MAX_PARTICIPANTS = 40
TOC_BOOKMARK_BASE = 239877506


def _bookmark_name(index: int) -> str:
    return f"_Toc{TOC_BOOKMARK_BASE + index}"


def _add_bookmark(paragraph, index: int) -> None:
    bookmark_id = str(TOC_BOOKMARK_BASE + index)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), _bookmark_name(index))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_field_start(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend((begin, instr, separate))


def _add_field_end(paragraph) -> None:
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _ensure_toc_styles(document: Document) -> None:
    names = {style.name for style in document.styles}
    for name, style_id, left_twips in (("toc 1", "TOC1", None), ("toc 2", "TOC2", "220")):
        if name in names:
            continue
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style._element.set(qn("w:styleId"), style_id)
        style._element.attrib.pop(qn("w:customStyle"), None)
        for child in list(style._element):
            style._element.remove(child)
        for tag, attributes in (
            ("name", {"val": name}),
            ("basedOn", {"val": "Normal"}),
            ("next", {"val": "Normal"}),
            ("autoRedefine", {}),
            ("uiPriority", {"val": "39"}),
            ("unhideWhenUsed", {}),
        ):
            element = OxmlElement(f"w:{tag}")
            for key, value in attributes.items():
                element.set(qn(f"w:{key}"), value)
            style._element.append(element)
        properties = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "100")
        properties.append(spacing)
        if left_twips is not None:
            indent = OxmlElement("w:ind")
            indent.set(qn("w:left"), left_twips)
            properties.append(indent)
        style._element.append(properties)

    if "Hyperlink" not in names:
        style = document.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        style._element.attrib.pop(qn("w:customStyle"), None)
        for child in list(style._element):
            style._element.remove(child)
        for tag, attributes in (
            ("name", {"val": "Hyperlink"}),
            ("basedOn", {"val": "DefaultParagraphFont"}),
            ("uiPriority", {"val": "99"}),
            ("unhideWhenUsed", {}),
        ):
            element = OxmlElement(f"w:{tag}")
            for key, value in attributes.items():
                element.set(qn(f"w:{key}"), value)
            style._element.append(element)
        run_properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0000FF")
        color.set(qn("w:themeColor"), "hyperlink")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        run_properties.extend((color, underline))
        style._element.append(run_properties)


def _set_style_run_properties(style, children: list[tuple[str, dict[str, str]]]) -> None:
    properties = style._element.get_or_add_rPr()
    for child in list(properties):
        properties.remove(child)
    for tag, attributes in children:
        element = OxmlElement(f"w:{tag}")
        for key, value in attributes.items():
            element.set(qn(f"w:{key}"), value)
        properties.append(element)


def _align_styles_with_current_document(document: Document) -> None:
    themes = {
        "asciiTheme": "majorHAnsi",
        "eastAsiaTheme": "majorEastAsia",
        "hAnsiTheme": "majorHAnsi",
        "cstheme": "majorBidi",
    }
    normal = document.styles["Normal"]
    _set_style_run_properties(
        normal,
        [("rFonts", {"ascii": "Aptos", "hAnsi": "Aptos"}), ("color", {"val": "172131"})],
    )
    normal_ppr = normal._element.get_or_add_pPr()
    for child in list(normal_ppr):
        normal_ppr.remove(child)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "140")
    normal_ppr.append(spacing)

    _set_style_run_properties(
        document.styles["Title"],
        [
            ("rFonts", themes),
            ("b", {}),
            ("spacing", {"val": "5"}),
            ("kern", {"val": "28"}),
            ("sz", {"val": "56"}),
            ("szCs", {"val": "52"}),
        ],
    )
    _set_style_run_properties(
        document.styles["Heading 1"],
        [
            ("rFonts", themes),
            ("b", {}),
            ("bCs", {}),
            ("color", {"val": "174A84"}),
            ("sz", {"val": "40"}),
            ("szCs", {"val": "28"}),
        ],
    )
    _set_style_run_properties(
        document.styles["Heading 2"],
        [
            ("rFonts", themes),
            ("b", {}),
            ("bCs", {}),
            ("sz", {"val": "30"}),
            ("szCs", {"val": "26"}),
        ],
    )
    _set_style_run_properties(
        document.styles["Heading 3"],
        [
            ("rFonts", themes),
            ("b", {}),
            ("bCs", {}),
            ("sz", {"val": "24"}),
        ],
    )


def _toc_run(*, hidden: bool = False):
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    properties.append(OxmlElement("w:noProof"))
    if hidden:
        properties.append(OxmlElement("w:webHidden"))
    run.append(properties)
    return run


def _add_internal_hyperlink(paragraph, text: str, anchor: str, page: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_stop = OxmlElement("w:tab")
    tab_stop.set(qn("w:val"), "right")
    tab_stop.set(qn("w:leader"), "dot")
    tab_stop.set(qn("w:pos"), "9062")
    tabs.append(tab_stop)
    properties.append(tabs)
    paragraph_run_properties = OxmlElement("w:rPr")
    paragraph_run_properties.append(OxmlElement("w:noProof"))
    properties.append(paragraph_run_properties)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    properties.append(OxmlElement("w:noProof"))
    run.append(properties)
    content = OxmlElement("w:t")
    content.text = text
    run.append(content)
    hyperlink.append(run)

    tab_run = _toc_run(hidden=True)
    tab_run.append(OxmlElement("w:tab"))
    hyperlink.append(tab_run)

    begin_run = _toc_run(hidden=True)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    hyperlink.append(begin_run)

    instruction_run = _toc_run(hidden=True)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" PAGEREF {anchor} \\h "
    instruction_run.append(instruction)
    hyperlink.append(instruction_run)

    separator_run = _toc_run(hidden=True)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    separator_run.append(separator)
    hyperlink.append(separator_run)

    page_run = _toc_run(hidden=True)
    page_text = OxmlElement("w:t")
    page_text.text = str(page)
    page_run.append(page_text)
    hyperlink.append(page_run)

    end_run = _toc_run(hidden=True)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    hyperlink.append(end_run)
    paragraph._p.append(hyperlink)


def _add_toc(document: Document, model: RegulationDocument) -> None:
    _ensure_toc_styles(document)
    heading = document.add_heading("Spis treści", level=1)
    heading.paragraph_format.space_after = Pt(16)
    _add_bookmark(heading, 0)
    entries = [
        ("Spis treści", 2, 1),
        ("Konstrukcja regulaminu", 3, 1),
        *((chapter.title, index + 4, 1) for index, chapter in enumerate(model.chapters)),
        ("Historia wersji", 11, 1),
    ]
    for index, (title, page, level) in enumerate(entries):
        paragraph = document.add_paragraph(style=f"toc {level}")
        if index == 0:
            _add_field_start(paragraph, 'TOC \\o "1-2" \\h \\z \\u')
        _add_internal_hyperlink(paragraph, title, _bookmark_name(index), page)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(16)
    _add_field_end(spacer)
    if model.metadata["toc_note"]:
        note = document.add_paragraph(style="Tekst roboczy")
        note.add_run(model.metadata["toc_note"])
    document.add_page_break()


def _add_outline(document: Document, model: RegulationDocument) -> None:
    heading = document.add_heading("Konstrukcja regulaminu", level=1)
    _add_bookmark(heading, 1)
    document.add_paragraph(model.metadata["outline_intro"])
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(1.8), Cm(5.2), Cm(9.0))
    for column, width in zip(table.columns, widths):
        column.width = width
    set_table_borders(table, color=LINE, size="4")
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, width, text in zip(header.cells, widths, ("Rozdział", "Tytuł", "Zakres")):
        cell.width = width
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for index, chapter in enumerate(model.chapters, start=1):
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell)
        for cell, width, text in zip(cells, widths, (str(index), chapter.title, chapter.scope)):
            cell.width = width
            cell.text = text
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(BLUE)
        cells[1].paragraphs[0].runs[0].bold = True
        cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                run.font.name = "Aptos"
                run.font.size = Pt(9)
    document.add_page_break()


def _set_row_indivisible(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _keep_table_together(table) -> None:
    for row_index, row in enumerate(table.rows):
        _set_row_indivisible(row)
        keep_with_next = row_index < len(table.rows) - 1
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = keep_with_next


def _normalize_cell_margins_for_word(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            properties = cell._tc.get_or_add_tcPr()
            margins = properties.find(qn("w:tcMar"))
            if margins is None:
                continue
            for child in margins:
                if child.tag == qn("w:start"):
                    child.tag = qn("w:left")
                elif child.tag == qn("w:end"):
                    child.tag = qn("w:right")


def _add_rank_coefficients(document: Document, block: TableBlock):
    table = document.add_table(rows=1, cols=len(block.rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    widths = (Cm(2.2), Cm(11.2), Cm(3.2))
    for column, width in zip(table.columns, widths):
        column.width = width
    for row_index, values in enumerate(block.rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for index, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, "173F73")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EAF2FB")
            paragraph = cell.paragraphs[0]
            if row_index > 0:
                paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index != 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(text)
            run.font.size = Pt(9)
            if row_index == 0 or index in (0, 2):
                run.bold = True
            if row_index == 0:
                run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    _keep_table_together(table)
    return table


def _add_source_table(document: Document, model: RegulationDocument, block: TableBlock):
    """Native Word cells: timeline stages are columns, ordinary tables keep rows."""
    timeline = block.name.startswith("timeline-")
    rows = list(zip(*block.rows[1:])) if timeline else block.rows
    width = 16 / len(rows[0])
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column in table.columns:
        column.width = Cm(width)
    colors = ("173F73", "875600", "176342", "455468")
    for number, values in enumerate(rows):
        cells = table.add_row().cells
        for column, (cell, text) in enumerate(zip(cells, values)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            if timeline:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(resolve_references(model, text))
            run.font.size = Pt(9)
            run.bold = number == 0 or (timeline and number == 1)
            if number == 0:
                set_cell_shading(cell, colors[column] if timeline else colors[0])
                run.font.color.rgb = RGBColor(255, 255, 255)
            elif timeline or number % 2 == 0:
                set_cell_shading(cell, "EAF2FB")
    set_repeat_table_header(table.rows[0])
    _keep_table_together(table)
    return table


def _ensure_ztp_styles(document: Document) -> None:
    names = {style.name for style in document.styles}
    if "Tytuł paragrafu" not in names:
        style = document.styles.add_style("Tytuł paragrafu", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Heading 3"]
    if "Etykieta brudnopisu" not in names:
        style = document.styles.add_style("Etykieta brudnopisu", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.bold = True
        style.font.size = Pt(8)
        style.font.color.rgb = RGBColor.from_string("595959")


def _unit_prefix(unit: ZtpUnit, index: int) -> str:
    if unit.kind == "paragraph":
        return ""
    if unit.kind == "ust":
        return f"{index}. "
    if unit.kind == "pkt":
        return f"{index}) "
    if unit.kind == "lit":
        return f"{chr(96 + index)}) "
    if unit.kind == "tiret":
        return "– "
    if unit.kind == "double-tiret":
        return "–– "
    raise ValueError(f"Nieobsługiwany rodzaj jednostki: {unit.kind}")


def _add_unit(
    document: Document,
    model: RegulationDocument,
    unit: ZtpUnit,
    index: int,
    *,
    keep_together: bool = False,
) -> None:
    content = document.add_paragraph(style="Normal")
    if keep_together:
        content.paragraph_format.keep_together = True
    prefix = _unit_prefix(unit, index)
    if prefix:
        content.add_run(prefix)
    for inline_run in unit.runs:
        run = content.add_run(resolve_references(model, inline_run.text))
        if inline_run.bold:
            run.bold = True
    indents = {
        "paragraph": (0.15, 0.0),
        "ust": (0.55, -0.4),
        "pkt": (1.05, -0.45),
        "lit": (1.55, -0.45),
        "tiret": (2.05, -0.45),
        "double-tiret": (2.55, -0.55),
    }
    left, first = indents[unit.kind]
    content.paragraph_format.left_indent = Cm(left)
    if first:
        content.paragraph_format.first_line_indent = Cm(first)
    if unit.status != "accepted":
        for run in content.runs:
            run.font.color.rgb = RGBColor.from_string("595959")
    if unit.identifier == "wspolczynniki-wprowadzenie":
        content.paragraph_format.keep_with_next = True
    for child_index, child in enumerate(unit.children, start=1):
        _add_unit(document, model, child, child_index)


def _add_status_label(document: Document, status: str) -> None:
    text = {
        "source-draft": "BRUDNOPIS ZE ŹRÓDŁA — DO OPRACOWANIA",
        "placeholder": "MIEJSCE DO UZUPEŁNIENIA — WYMAGA DECYZJI",
    }[status]
    label = document.add_paragraph(text, style="Etykieta brudnopisu")
    label.paragraph_format.keep_with_next = True


def _add_external_link(paragraph, url: str) -> None:
    parsed = urlsplit(url)
    if parsed.scheme != "https" or not parsed.netloc:
        raise ValueError("Nieprawidłowy adres HTTPS materiału pomocniczego")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.append(size)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _add_resource_links(document: Document, model: RegulationDocument) -> None:
    if "resource_links_title" not in model.metadata:
        return
    heading = document.add_paragraph(style="Normal")
    heading.add_run(model.metadata["resource_links_title"]).bold = True
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.keep_with_next = True
    note = document.add_paragraph(model.metadata["resource_links_note"], style="Normal")
    note.paragraph_format.keep_with_next = True
    for index, (label, key) in enumerate((
        ("Ranking", "resource_ranking_url"),
        ("Tabela punktacji", "resource_table_url"),
        ("Kalkulator punktów", "resource_calculator_url"),
    )):
        url = model.metadata[key]
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.keep_with_next = index < 2
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run(label + ": ").bold = True
        _add_external_link(paragraph, url)


def _add_chapters(document: Document, model: RegulationDocument) -> None:
    _ensure_ztp_styles(document)
    section_number = 0
    for chapter_index, chapter in enumerate(model.chapters, start=1):
        label = document.add_paragraph(f"Rozdział {chapter_index}", style="Etykieta rozdzialu")
        label.paragraph_format.page_break_before = True
        heading = document.add_paragraph(chapter.title, style="Heading 1")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_bookmark(heading, chapter_index + 1)
        for section in chapter.sections:
            section_number += 1
            paragraph = document.add_paragraph(f"§ {section_number}", style="Paragraf")
            paragraph.paragraph_format.keep_with_next = True
            if section.identifier in {"cel", "definicje"}:
                paragraph.paragraph_format.left_indent = Cm(0.15)
            if section.title:
                title = document.add_paragraph(section.title, style="Tytuł paragrafu")
                title.paragraph_format.keep_with_next = True
            active_status = "accepted"
            unit_index = 0
            for block in section.blocks:
                if isinstance(block, ZtpUnit):
                    unit_index += 1
                    if block.status != "accepted" and block.status != active_status:
                        _add_status_label(document, block.status)
                    active_status = block.status
                    _add_unit(
                        document,
                        model,
                        block,
                        unit_index,
                        keep_together=unit_index == 1,
                    )
                elif block.name == "rank-coefficients":
                    active_status = "accepted"
                    _add_rank_coefficients(document, block)
                else:
                    active_status = "accepted"
                    _add_source_table(document, model, block)
            if section.identifier == "publikacja-rankingu":
                _add_resource_links(document, model)


def _is_power_of_two(value: int) -> bool:
    return value > 0 and (value & (value - 1)) == 0


def _score(place: int, participants: int) -> float:
    base = min(50.0, 10.0 * math.log2(max(2, participants)))
    place_points = base - (base - 1.0) * math.log(place) / math.log(participants)
    won_rounds = max(
        0,
        math.floor(math.log2(participants))
        - math.ceil(math.log2(place))
        + (0 if _is_power_of_two(participants) else 1),
    )
    podium_factor = {1: 3, 2: 2, 3: 1}.get(place, 0)
    podium = podium_factor * 3.0 * participants ** (1.0 / 3.0)
    return place_points + won_rounds * 10.0 + podium


def _format_points(value: float) -> str:
    return f"{value:.1f}".replace(".", ",")


def _format_annex_cell(cell, text: str, *, header=False, first_column=False, alternate=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        set_cell_shading(cell, "173F73")
    elif first_column:
        set_cell_shading(cell, "EAF2FB")
    elif alternate:
        set_cell_shading(cell, "F6F8FB")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.size = Pt(7.5)
    run.bold = header or first_column
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)


def _add_points_annex(document: Document, model: RegulationDocument) -> None:
    document.add_page_break()
    label = document.add_paragraph("ZAŁĄCZNIK NR 1", style="Etykieta rozdzialu")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = document.add_paragraph(
        "Tabela punktacji Pucharu Polski Weteranów w szermierce",
        style="Heading 1",
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(style="Normal")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Punkty rankingowe za miejsce zdobyte w stawce liczącej od 4 do 40 zawodników "
        "— sezon 2026/2027"
    ).italic = True
    coefficient = document.add_paragraph(style="Normal")
    coefficient.alignment = WD_ALIGN_PARAGRAPH.CENTER
    coefficient.add_run("Współczynnik rangi PPW: 1,0").bold = True

    for first_place, last_place in PLACE_BLOCKS:
        places = list(range(first_place, last_place + 1))
        first_participant = max(MIN_PARTICIPANTS, first_place)
        participants = list(range(first_participant, MAX_PARTICIPANTS + 1))
        for start in range(0, len(participants), PARTICIPANT_CHUNK_SIZE):
            chunk = participants[start : start + PARTICIPANT_CHUNK_SIZE]
            heading = document.add_paragraph(
                f"Miejsca {first_place}–{last_place} · stawka {chunk[0]}–{chunk[-1]} zawodników",
                style="Heading 2",
            )
            heading.paragraph_format.keep_with_next = True
            table = document.add_table(rows=1, cols=1 + len(places))
            table.style = "Table Grid"
            table.autofit = False
            header = table.rows[0]
            set_repeat_table_header(header)
            header.cells[0].width = Cm(3.0)
            _format_annex_cell(header.cells[0], "Liczba zawodników w stawce", header=True)
            for index, place in enumerate(places, start=1):
                header.cells[index].width = Cm(1.25)
                _format_annex_cell(header.cells[index], str(place), header=True)
            for row_number, participant_count in enumerate(chunk):
                row = table.add_row()
                row.cells[0].width = Cm(3.0)
                _format_annex_cell(row.cells[0], str(participant_count), first_column=True)
                for column_index, place in enumerate(places, start=1):
                    row.cells[column_index].width = Cm(1.25)
                    value = _format_points(_score(place, participant_count)) if place <= participant_count else "—"
                    _format_annex_cell(
                        row.cells[column_index],
                        value,
                        alternate=row_number % 2 == 1,
                    )
            _keep_table_together(table)
    note = document.add_paragraph(
        "Pełna tabela dla stawek liczących od 4 do 300 zawodników oraz kalkulator punktów "
        "są publikowane przez SPWS na stronie internetowej."
    )
    note.paragraph_format.space_before = Pt(8)
    if "resource_table_url" in model.metadata:
        note.paragraph_format.keep_with_next = True
        link = document.add_paragraph(style="Normal")
        link.paragraph_format.keep_together = True
        link.add_run("Pełna tabela punktacji — wersja nieoficjalna: ").bold = True
        _add_external_link(link, model.metadata["resource_table_url"])
    document.add_page_break()


def _add_history(document: Document, model: RegulationDocument) -> None:
    history_heading = document.add_heading("Historia wersji", level=1)
    _add_bookmark(history_heading, 9)
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=LINE, size="4")
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, ("Wersja", "Data", "Zakres zmian", "Status")):
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    cells = table.add_row().cells
    values = (
        model.metadata["version"],
        model.metadata["project_date"],
        model.metadata["history_scope"],
        model.metadata["status"],
    )
    for cell, text in zip(cells, values):
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.font.name = "Aptos"
        run.font.size = Pt(9)
    _keep_table_together(table)


def _update_fields_on_open(document: Document) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _normalize_empty_first_page_header(document: Document) -> None:
    paragraph = document.sections[0].first_page_header.paragraphs[0]
    for run in paragraph._p.xpath("w:r"):
        paragraph._p.remove(run)
    properties = paragraph._p.get_or_add_pPr()
    run_properties = properties.find(qn("w:rPr"))
    if run_properties is None:
        run_properties = OxmlElement("w:rPr")
        properties.append(run_properties)
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.append(fonts)
    fonts.set(qn("w:hint"), "eastAsia")


def _normalize_cover_run_properties(document: Document) -> None:
    document.paragraphs[0].runs[0].font.size = None
    cover_table = document.tables[0]
    for row in cover_table.rows[1:]:
        row.cells[1].paragraphs[0].runs[0].bold = None


def _ensure_blank_paragraph_after_tables(document: Document) -> None:
    for table in document.tables:
        following = table._tbl.getnext()
        if following is not None and following.tag == qn("w:p"):
            text = "".join(element.text or "" for element in following.iter(qn("w:t")))
            if not text:
                continue
        separator = document.add_paragraph(style="Normal")
        table._tbl.addnext(separator._p)


def build_document(source: Path, output: Path) -> Path:
    source = source.resolve()
    output = output.resolve()
    if output == source or output.suffix.lower() != ".docx":
        raise ValueError("Wyjście musi być osobnym plikiem DOCX")
    model = parse_regulation_source(source)
    document = Document()
    configure_styles(document)
    _align_styles_with_current_document(document)
    for section in document.sections:
        configure_page(section)
    add_header_footer(document.sections[0])
    add_cover(document)
    _normalize_empty_first_page_header(document)
    _normalize_cover_run_properties(document)
    _add_toc(document, model)
    _add_outline(document, model)
    _add_chapters(document, model)
    if model.has_points_annex:
        _add_points_annex(document, model)
    _add_history(document, model)
    _ensure_blank_paragraph_after_tables(document)
    for table in document.tables:
        _keep_table_together(table)
        _normalize_cell_margins_for_word(table)
    _update_fields_on_open(document)
    document.core_properties.title = model.metadata["title"]
    document.core_properties.subject = model.metadata["subject"]
    document.core_properties.author = "Komisja regulaminowa SPWS"
    document.core_properties.last_modified_by = "Komisja regulaminowa SPWS"
    document.core_properties.comments = model.metadata["comments"]
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="Buduje Regulamin DOCX z warstwy treści Markdown")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    print(build_document(arguments.source, arguments.output))


if __name__ == "__main__":
    main()
