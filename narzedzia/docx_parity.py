from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
CORE_TIME_NS = {"dcterms": "http://purl.org/dc/terms/"}
TRACKED_STYLES = (
    "Normal",
    "Title",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Etykieta rozdzialu",
    "Paragraf",
    "Tekst roboczy",
    "toc 1",
    "toc 2",
)
RENDER_STYLE_IDS = (
    "Normal",
    "Title",
    "Heading1",
    "Heading2",
    "Heading3",
    "TOC1",
    "TOC2",
    "Hyperlink",
)


def _length(value) -> int | None:
    return None if value is None else int(value)


def _color(value) -> str | None:
    return None if value is None else str(value)


def _paragraph_layout(paragraph) -> dict[str, object]:
    formatting = paragraph.paragraph_format
    return {
        "style": paragraph.style.name,
        "alignment": None if paragraph.alignment is None else int(paragraph.alignment),
        "page_break_before": formatting.page_break_before,
        "keep_with_next": formatting.keep_with_next,
        "space_before": _length(formatting.space_before),
        "space_after": _length(formatting.space_after),
        "left_indent": _length(formatting.left_indent),
        "first_line_indent": _length(formatting.first_line_indent),
        "line_spacing": None if formatting.line_spacing is None else str(formatting.line_spacing),
    }


def _run_layout(run) -> dict[str, object]:
    return {
        "text": run.text,
        "size": _length(run.font.size),
        "bold": bool(run.bold),
        "italic": bool(run.italic),
        "color": _color(run.font.color.rgb),
    }


def _normalized_runs(paragraph) -> list[dict[str, object]]:
    groups: list[dict[str, object]] = []
    for run in paragraph.runs:
        current = _run_layout(run)
        formatting = {key: value for key, value in current.items() if key != "text"}
        if groups and all(groups[-1][key] == value for key, value in formatting.items()):
            groups[-1]["text"] += current["text"]
        else:
            groups.append(current)
    return groups


def _inline_emphasis_contract(document: Document) -> list[dict[str, object]]:
    def emphasized_segments(paragraph, attribute: str) -> list[str]:
        segments: list[str] = []
        active = ""
        for run in paragraph.runs:
            if bool(getattr(run, attribute)):
                active += run.text
            elif active:
                segments.append(active)
                active = ""
        if active:
            segments.append(active)
        return segments

    result = []
    for paragraph in document.paragraphs:
        bold = emphasized_segments(paragraph, "bold")
        italic = emphasized_segments(paragraph, "italic")
        if bold or italic:
            result.append({"text": paragraph.text, "bold": bold, "italic": italic})
    return result


def _cell_layout(cell) -> dict[str, object]:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    return {
        "width": _length(cell.width),
        "vertical_alignment": (
            None if cell.vertical_alignment is None else int(cell.vertical_alignment)
        ),
        "shading": None if shading is None else shading.get(qn("w:fill")),
        "paragraphs": [
            {
                "text": paragraph.text,
                "layout": _paragraph_layout(paragraph),
                "runs": _normalized_runs(paragraph),
            }
            for paragraph in cell.paragraphs
        ],
    }


def _normalized_xml_element(element) -> tuple | None:
    if element is None:
        return None
    if etree.QName(element).localname == "rsid":
        return None
    attributes = tuple(
        sorted(
            (etree.QName(name).localname, value)
            for name, value in element.attrib.items()
            if etree.QName(name).localname not in {"rsid", "rsidR", "rsidRPr"}
        )
    )
    return (
        etree.QName(element).localname,
        attributes,
        (element.text or "").strip(),
        tuple(
            normalized
            for child in element
            if (normalized := _normalized_xml_element(child)) is not None
        ),
    )


def _render_style_contract(styles_root) -> dict[str, tuple | None]:
    return {
        style_id: _normalized_xml_element(
            next(
                iter(
                    styles_root.xpath(
                        f"w:style[@w:styleId='{style_id}']",
                        namespaces=WORD_NS,
                    )
                ),
                None,
            )
        )
        for style_id in RENDER_STYLE_IDS
    }


def _toc_hyperlink_contract(main_root) -> list[dict[str, object]]:
    paragraphs = main_root.xpath(
        "//w:body/w:p[w:pPr/w:pStyle[@w:val='TOC1' or @w:val='TOC2']]",
        namespaces=WORD_NS,
    )
    return [
        {
            "style": paragraph.xpath("string(w:pPr/w:pStyle/@w:val)", namespaces=WORD_NS),
            "tabs": [
                tuple(sorted((etree.QName(name).localname, value) for name, value in tab.attrib.items()))
                for tab in paragraph.xpath("w:pPr/w:tabs/w:tab", namespaces=WORD_NS)
            ],
            "hyperlinks": [
                {
                    "anchor": hyperlink.get(qn("w:anchor")),
                    "history": hyperlink.get(qn("w:history")),
                    "text": "".join(hyperlink.xpath(".//w:t/text()", namespaces=WORD_NS)),
                    "field_instructions": hyperlink.xpath(
                        ".//w:instrText/text()", namespaces=WORD_NS
                    ),
                    "run_styles": hyperlink.xpath(
                        ".//w:rStyle/@w:val", namespaces=WORD_NS
                    ),
                }
                for hyperlink in paragraph.xpath("w:hyperlink", namespaces=WORD_NS)
            ],
        }
        for paragraph in paragraphs
    ]


def _toc_field_boundary_contract(main_root) -> dict[str, int]:
    paragraphs = main_root.xpath(
        "//w:body/w:p[w:pPr/w:pStyle[@w:val='TOC1' or @w:val='TOC2']]",
        namespaces=WORD_NS,
    )
    last_entry = paragraphs[-1]
    following = last_entry.getnext()
    selector = "w:r/w:fldChar[@w:fldCharType='end']"
    return {
        "last_entry_direct_end": len(last_entry.xpath(selector, namespaces=WORD_NS)),
        "following_paragraph_direct_end": len(following.xpath(selector, namespaces=WORD_NS)),
    }


def document_content_contract(path: Path) -> dict[str, object]:
    document = Document(path)
    paragraphs = list(document.paragraphs)
    tables = list(document.tables)
    paragraph_by_element = {id(paragraph._p): paragraph for paragraph in paragraphs}
    table_by_element = {id(table._tbl): table for table in tables}
    blocks: list[dict[str, object]] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = paragraph_by_element[id(child)]
            blocks.append(
                {
                    "kind": "paragraph",
                    "style": paragraph.style.name,
                    "text": paragraph.text,
                }
            )
        elif child.tag == qn("w:tbl"):
            table = table_by_element[id(child)]
            blocks.append(
                {
                    "kind": "table",
                    "rows": [[cell.text for cell in row.cells] for row in table.rows],
                }
            )
    return {
        "blocks": blocks,
        "headers": [
            paragraph.text
            for section in document.sections
            for paragraph in section.header.paragraphs
        ],
        "footers": [
            paragraph.text
            for section in document.sections
            for paragraph in section.footer.paragraphs
        ],
    }


def document_layout_contract(path: Path) -> dict[str, object]:
    document = Document(path)
    core = document.core_properties
    with ZipFile(path) as archive:
        field_parts = []
        for name in sorted(
            item for item in archive.namelist() if item.startswith("word/") and item.endswith(".xml")
        ):
            root = etree.fromstring(archive.read(name))
            instructions = [
                instruction.text or ""
                for instruction in root.xpath("//w:instrText", namespaces=WORD_NS)
            ]
            if instructions:
                field_parts.append((name, instructions))
        main = etree.fromstring(archive.read("word/document.xml"))
        styles = etree.fromstring(archive.read("word/styles.xml"))
        bookmarks = [
            element.get(qn("w:name"))
            for element in main.xpath("//w:bookmarkStart", namespaces=WORD_NS)
            if (element.get(qn("w:name")) or "").startswith("_Toc")
        ]

    return {
        "paragraphs": [_paragraph_layout(paragraph) for paragraph in document.paragraphs],
        "inline_emphasis": _inline_emphasis_contract(document),
        "tables": [
            {
                "style": None if table.style is None else table.style.name,
                "alignment": None if table.alignment is None else int(table.alignment),
                "autofit": table.autofit,
                "grid_widths": [
                    column.get(qn("w:w"))
                    for column in table._tbl.tblGrid.findall(qn("w:gridCol"))
                ],
                "rows": [
                    {
                        "cant_split": row._tr.get_or_add_trPr().find(qn("w:cantSplit"))
                        is not None,
                        "cells": [_cell_layout(cell) for cell in row.cells],
                    }
                    for row in table.rows
                ],
            }
            for table in document.tables
        ],
        "sections": [
            {
                "page_width": _length(section.page_width),
                "page_height": _length(section.page_height),
                "top_margin": _length(section.top_margin),
                "bottom_margin": _length(section.bottom_margin),
                "left_margin": _length(section.left_margin),
                "right_margin": _length(section.right_margin),
                "header_distance": _length(section.header_distance),
                "footer_distance": _length(section.footer_distance),
                "different_first_page": section.different_first_page_header_footer,
            }
            for section in document.sections
        ],
        "required_styles": [name for name in TRACKED_STYLES if name in document.styles],
        "render_styles": _render_style_contract(styles),
        "toc_hyperlinks": _toc_hyperlink_contract(main),
        "toc_field_boundary": _toc_field_boundary_contract(main),
        "first_page_header_direct_runs": [
            len(paragraph._p.xpath("w:r"))
            for section in document.sections
            for paragraph in section.first_page_header.paragraphs
        ],
        "first_page_header_paragraph_properties": [
            _normalized_xml_element(paragraph._p.pPr)
            for section in document.sections
            for paragraph in section.first_page_header.paragraphs
        ],
        "fields": field_parts,
        "toc_bookmarks": bookmarks,
        "core_properties": {
            "title": core.title,
            "subject": core.subject,
            "author": core.author,
            "last_modified_by": core.last_modified_by,
            "comments": core.comments,
        },
    }


def normalized_package_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            payload = archive.read(name)
            if name == "docProps/core.xml":
                root = etree.fromstring(payload)
                for element in root.xpath(
                    "//dcterms:created | //dcterms:modified",
                    namespaces=CORE_TIME_NS,
                ):
                    element.text = "NORMALIZED"
                payload = etree.tostring(root)
            digest.update(name.encode("utf-8"))
            digest.update(b"\0")
            digest.update(payload)
            digest.update(b"\0")
    return digest.hexdigest()


def parity_report(reference: Path, candidate: Path) -> dict[str, object]:
    content_match = document_content_contract(reference) == document_content_contract(candidate)
    layout_match = document_layout_contract(reference) == document_layout_contract(candidate)
    return {
        "reference": str(reference),
        "candidate": str(candidate),
        "content_match": content_match,
        "layout_match": layout_match,
        "passed": content_match and layout_match,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Porównuje wzorcowy i wygenerowany DOCX")
    parser.add_argument("reference", type=Path)
    parser.add_argument("candidate", type=Path)
    arguments = parser.parse_args()
    report = parity_report(arguments.reference, arguments.candidate)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(0 if report["passed"] else 1)


if __name__ == "__main__":
    main()
