from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-reprezentacji-Polski-weteranów-w-szermierce_2026.docx"
)

BLUE = "174A84"
DARK = "172131"
MUTED = "667386"
LIGHT = "E7EBF1"
LINE = "CDD4DE"
RED = "A03A2D"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_bottom_border(paragraph, color=LINE, size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.styles["Title"]
    title.font.name = "Georgia"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK)
    title.paragraph_format.space_after = Pt(14)

    for name, size, color in (
        ("Heading 1", 20, BLUE),
        ("Heading 2", 15, DARK),
        ("Heading 3", 12, DARK),
    ):
        style = doc.styles[name]
        style.font.name = "Georgia" if name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(8)

    chapter = doc.styles.add_style("Rozdzial", WD_STYLE_TYPE.PARAGRAPH)
    chapter.base_style = doc.styles["Heading 1"]
    chapter.font.name = "Georgia"
    chapter.font.size = Pt(20)
    chapter.font.bold = True
    chapter.font.color.rgb = RGBColor.from_string(BLUE)
    chapter.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chapter.paragraph_format.space_before = Pt(0)
    chapter.paragraph_format.space_after = Pt(20)
    chapter.paragraph_format.keep_with_next = True
    chapter._element.get_or_add_pPr().append(OxmlElement("w:pageBreakBefore"))

    section_label = doc.styles.add_style("Etykieta rozdzialu", WD_STYLE_TYPE.PARAGRAPH)
    section_label.font.name = "Aptos"
    section_label.font.size = Pt(9)
    section_label.font.bold = True
    section_label.font.color.rgb = RGBColor.from_string(BLUE)
    section_label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section_label.paragraph_format.space_after = Pt(5)
    section_label.paragraph_format.keep_with_next = True

    paragraph_mark = doc.styles.add_style("Paragraf", WD_STYLE_TYPE.PARAGRAPH)
    paragraph_mark.font.name = "Georgia"
    paragraph_mark.font.size = Pt(13)
    paragraph_mark.font.bold = True
    paragraph_mark.font.color.rgb = RGBColor.from_string(DARK)
    paragraph_mark.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_mark.paragraph_format.space_before = Pt(15)
    paragraph_mark.paragraph_format.space_after = Pt(8)
    paragraph_mark.paragraph_format.keep_with_next = True

    placeholder = doc.styles.add_style("Tekst roboczy", WD_STYLE_TYPE.PARAGRAPH)
    placeholder.base_style = normal
    placeholder.font.italic = True
    placeholder.font.color.rgb = RGBColor.from_string(MUTED)


def configure_page(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_bottom_border(p)
    run = p.add_run("Regulamin powoływania reprezentacji Polski weteranów")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Projekt · wersja 0.1     |     Strona ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_field(p, "PAGE", "1")
    p.add_run(" z ")
    add_field(p, "NUMPAGES", "1")


def add_cover(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    footer = section.first_page_footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Projekt regulaminu  ·  sezon 2026/2027")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    project = doc.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    project.paragraph_format.space_after = Pt(58)
    run = project.add_run("PROJEKT")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(RED)
    r_pr = run._r.get_or_add_rPr()
    border = OxmlElement("w:bdr")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "14")
    border.set(qn("w:space"), "4")
    border.set(qn("w:color"), RED)
    r_pr.append(border)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(32)
    # Keep the existing cover spacing, without a decorative line above the title.

    title = doc.add_paragraph(style="Title")
    title.add_run("Regulamin powoływania reprezentacji Polski weteranów w szermierce")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(96)
    run = subtitle.add_run("w sezonie 2026/2027")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(5.2)
    table.columns[1].width = Cm(10.0)
    set_table_borders(table, color=LINE, size="4")
    values = (
        ("Wersja dokumentu", "0.1"),
        ("Status", "projekt do konsultacji"),
        ("Data projektu", "[data]"),
    )
    for row, (label, value) in zip(table.rows, values):
        row.cells[0].width = Cm(5.2)
        row.cells[1].width = Cm(10.0)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], LIGHT)
        for cell in row.cells:
            set_cell_margins(cell)
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        run.font.name = "Aptos"
        run.font.size = Pt(10)
        run.bold = label == "Wersja dokumentu"

    doc.add_page_break()


def add_toc(doc):
    heading = doc.add_heading("Spis treści", level=1)
    heading.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "Spis treści zostanie zaktualizowany po otwarciu dokumentu w Wordzie.")
    p.paragraph_format.space_after = Pt(16)
    note = doc.add_paragraph(style="Tekst roboczy")
    note.add_run("W Wordzie wybierz spis treści i polecenie „Aktualizuj tabelę”, jeżeli nie odświeży się automatycznie.")
    doc.add_page_break()


def add_outline(doc):
    doc.add_heading("Konstrukcja regulaminu", level=1)
    intro = doc.add_paragraph()
    intro.add_run(
        "Poniższy układ odzwierciedla metodę wyłaniania reprezentacji: dwie główne składowe "
        "oraz dwa elementy wspierające. Treść normatywna zostanie dodana po zatwierdzeniu kolejnych decyzji."
    )
    chapters = (
        ("I", "Postanowienia ogólne", "cel, zakres, organy i podstawowe definicje"),
        ("II", "Ranking indywidualny", "wyniki, punktacja, klasyfikacja i publikacja"),
        ("III", "Powołania do startów indywidualnych", "uprawnienia, kolejność i rezygnacje"),
        ("IV", "Dobór składu drużyny", "pula kandydatów, kryteria, role i odpowiedzialność"),
        ("V", "Terminarz procesu powoływania", "zamknięcie danych, konsultacje i ogłoszenie nominacji"),
        ("VI", "Ocena regulaminu i doskonalenie metody", "ewaluacja sezonowa i zasady zmian"),
        ("VII", "Postanowienia końcowe", "wejście w życie, przepisy przejściowe i załączniki"),
    )
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=LINE, size="4")
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, ("Rozdział", "Tytuł", "Zakres roboczy")):
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for numeral, title, scope in chapters:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell)
        cells[0].text = numeral
        cells[1].text = title
        cells[2].text = scope
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(BLUE)
        cells[1].paragraphs[0].runs[0].bold = True
        cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                run.font.name = "Aptos"
                run.font.size = Pt(9)


def add_chapter(doc, numeral, title, paragraphs):
    doc.add_page_break()
    doc.add_paragraph(f"ROZDZIAŁ {numeral}", style="Etykieta rozdzialu")
    heading = doc.add_paragraph(title, style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for mark, lines in paragraphs:
        doc.add_paragraph(mark, style="Paragraf")
        for idx, text in enumerate(lines, 1):
            p = doc.add_paragraph(style="Tekst roboczy")
            p.paragraph_format.left_indent = Cm(0.15)
            run = p.add_run(f"{idx}. ")
            run.bold = True
            p.add_run(text)


def add_skeleton(doc):
    chapters = (
        ("I", "Postanowienia ogólne", (("§ 1", ("[Cel regulaminu.]", "[Zakres spraw regulowanych dokumentem.]")), ("§ 2", ("[Definicje pojęć używanych w regulaminie.]",)))),
        ("II", "Ranking indywidualny", (("§ 3", ("[Cel i znaczenie rankingu indywidualnego.]",)), ("§ 4", ("[Kategorie, zawody i wyniki uwzględniane w rankingu.]", "[Zasady punktacji i publikacji rankingu.]")))),
        ("III", "Powołania do startów indywidualnych", (("§ 5", ("[Zasady kwalifikacji i ustalania kolejności kandydatów.]", "[Rezygnacja, zastępstwo i sytuacje szczególne.]")),)),
        ("IV", "Dobór składu drużyny", (("§ 6", ("[Pula kandydatów do drużyny.]", "[Wymogi kategorii wiekowych i dostępności.]")), ("§ 7", ("[Kryteria wyboru składu, role i odpowiedzialność za decyzję.]",)))),
        ("V", "Terminarz procesu powoływania", (("§ 8", ("[Daty graniczne procesu.]", "[Deklaracje, weryfikacja danych i ogłoszenie decyzji.]")),)),
        ("VI", "Ocena regulaminu i doskonalenie metody", (("§ 9", ("[Zakres i termin oceny posezonowej.]", "[Zasady przygotowania zmian na kolejny sezon.]")),)),
        ("VII", "Postanowienia końcowe", (("§ 10", ("[Przepisy przejściowe.]", "[Wejście regulaminu w życie.]")),)),
    )
    for numeral, title, paragraphs in chapters:
        add_chapter(doc, numeral, title, paragraphs)


def add_version_history(doc):
    doc.add_page_break()
    doc.add_heading("Historia wersji", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=LINE, size="4")
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, ("Wersja", "Data", "Zakres zmian", "Status")):
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    cells = table.add_row().cells
    for cell, text in zip(cells, ("0.1", "[data]", "Pierwszy prototyp struktury i formatowania", "projekt do konsultacji")):
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = "Aptos"
        run.font.size = Pt(9)

    doc.add_heading("Informacja o prototypie", level=2)
    p = doc.add_paragraph(style="Tekst roboczy")
    p.add_run(
        "Dokument służy wyłącznie do oceny struktury i formatowania. Tekst w nawiasach kwadratowych "
        "nie stanowi projektu postanowień regulaminu."
    )


def update_fields_on_open(doc):
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
    add_header_footer(doc.sections[0])
    add_cover(doc)
    add_toc(doc)
    add_outline(doc)
    add_skeleton(doc)
    add_version_history(doc)
    update_fields_on_open(doc)
    doc.core_properties.title = "Regulamin powoływania reprezentacji Polski weteranów w szermierce"
    doc.core_properties.subject = "Prototyp struktury i formatowania — sezon 2026/2027"
    doc.core_properties.author = "Komisja regulaminowa"
    doc.core_properties.comments = "Projekt do konsultacji; treść normatywna nie została jeszcze uzgodniona."
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
