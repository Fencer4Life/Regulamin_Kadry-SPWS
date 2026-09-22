"""Coverage, reference reconstruction and isolated before/after E2E evidence."""
import argparse
import hashlib
import html
import json
import re
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

from docx import Document
from narzedzia.build_regulamin_docx import build_document
from narzedzia.decision_patch import apply_decision, format_fragments
from narzedzia.docx_parity import normalized_package_sha256


def visible_source(text):
    text = text.split('\n+++\n', 1)[1]
    text = re.sub(r'<!--.*?-->', '', text, flags=re.S)
    text = re.sub(r'\[(?:unit|section|chapter|status):[^]]+\]', '', text)
    text = re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', text)
    text = '\n'.join(line.replace('|', ' ') if line.startswith('|') else line for line in text.splitlines())
    return ' '.join(text.replace('**', '').split())


def coverage(source, reference):
    visible = visible_source(source.read_text(encoding='utf-8'))
    document = Document(reference)
    groups = {'body': [(f'paragraph[{i}]', p) for i, p in enumerate(document.paragraphs)],
              'tables': [(f'table[{i}]/row[{j}]/cell[{k}]/paragraph[{l}]', p)
                         for i, t in enumerate(document.tables) for j, row in enumerate(t.rows)
                         for k, cell in enumerate(row.cells) for l, p in enumerate(cell.paragraphs)],
              'running': [(f'section[{i}]/{name}/paragraph[{j}]', p)
                          for i, section in enumerate(document.sections)
                          for name in ('header', 'footer', 'first_page_header', 'first_page_footer')
                          for j, p in enumerate(getattr(section, name).paragraphs)]}
    result = {'groups': {}, 'missing': [], 'items': []}
    for name, paragraphs in groups.items():
        checked = found = 0
        for location, paragraph in paragraphs:
            text = ' '.join(paragraph.text.split())
            if not text:
                continue
            checked += 1
            offset = visible.find(text)
            found += offset >= 0
            item = {'location': location, 'text': text, 'visible_source_offset': offset}
            result['items'].append(item)
            if offset < 0:
                result['missing'].append(item)
        result['groups'][name] = {'checked': checked, 'found': found}
    return result


def package(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def prove(source, reference, output):
    source, reference, output = source.resolve(), reference.resolve(), output.resolve()
    if output == source.parent or output == reference.parent:
        raise ValueError('Dowody muszą powstać w osobnym katalogu')
    output.mkdir(parents=True, exist_ok=True)
    baseline, repeated = output/'baseline.docx', output/'repeated.docx'
    build_document(source, baseline)
    build_document(source, repeated)
    result = {'reference_sha256': hashlib.sha256(reference.read_bytes()).hexdigest(),
              'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
              'coverage': coverage(source, reference),
              'reference_package_equal': normalized_package_sha256(reference) == normalized_package_sha256(baseline),
              'repeated_build_equal': normalized_package_sha256(baseline) == normalized_package_sha256(repeated)}
    # Exactly one ordinary unit, unchanged test target across runs. No tracked file is edited.
    marker = '<!-- unit:publikacja-rankingu-zrodlo --> '
    original_line = next(line for line in source.read_text(encoding='utf-8').splitlines() if line.startswith(marker))
    old = original_line[len(marker):]
    new = old.removesuffix('.')+' [TEST E2E].'
    test_source, test_docx, card = output/'E2E_TEST.md', output/'E2E_TEST.docx', output/'E2E_TEST_DECISION.md'
    test_source.write_bytes(source.read_bytes())
    card.write_text('---\nstatus: przyjęta\n---\n'+format_fragments(original_line, marker+new), encoding='utf-8')
    apply_decision(card, test_source, test_docx)
    before, after = package(baseline), package(test_docx)
    changed = sorted(key for key in before.keys() | after.keys() if before.get(key) != after.get(key))
    old_xml, new_xml = escape(old).encode(), escape(new).encode()
    count = before['word/document.xml'].count(old_xml)
    expected = before['word/document.xml'].replace(old_xml, new_xml, 1)
    result.update(e2e_changed_parts=changed, e2e_occurrences=count,
                  e2e_only_expected_replacement=(count == 1 and after['word/document.xml'] == expected
                                                 and changed == ['word/document.xml']),
                  e2e_before=old, e2e_after=new,
                  e2e_paragraph_index=next(i for i, p in enumerate(Document(baseline).paragraphs) if p.text == old))
    result['passed'] = (not result['coverage']['missing'] and result['reference_package_equal']
                        and result['repeated_build_equal'] and result['e2e_only_expected_replacement'])
    (output/'evidence.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    groups = result['coverage']['groups']
    rows = ''.join(f'<tr><td>{name}</td><td>{g["found"]}/{g["checked"]}</td></tr>' for name,g in groups.items())
    report = ('<!doctype html><html lang="pl"><meta charset="utf-8"><title>Pokrycie źródła i E2E</title>'
              '<style>body{max-width:1000px;margin:40px auto;font:18px/1.6 system-ui;padding:24px}'
              'td,th{padding:12px;border:1px solid #bbc}pre{white-space:pre-wrap;overflow-wrap:anywhere}</style>'
              '<h1>Pokrycie treści DOCX w Markdown i E2E</h1><table><tr><th>Obszar</th><th>Znaleziono/sprawdzono</th></tr>'
              +rows+'</table><p>Pokrycie sprawdza obecność każdego niepustego akapitu i komórki. '
              'Kolejność i formatowanie kontroluje porównanie całych paczek DOCX. '
              'E2E dopuszcza wyłącznie jedną dosłowną zamianę w word/document.xml, bez przeniesienia węzła. '
              'Pliki E2E_TEST są wyłącznie materiałem testowym.</p><pre>'
              +html.escape(json.dumps({k:v for k,v in result.items() if k!='coverage'}, ensure_ascii=False, indent=2))+'</pre></html>')
    (output/'report.html').write_text(report, encoding='utf-8')
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('reference', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    result = prove(args.source, args.reference, args.output)
    print(json.dumps({**{k:v for k,v in result.items() if k!='coverage'},
                      'coverage': result['coverage']['groups']}, ensure_ascii=False, indent=2))
    if not result['passed']:
        raise SystemExit('FAILED: sprawdź evidence.json')


if __name__ == '__main__':
    main()
