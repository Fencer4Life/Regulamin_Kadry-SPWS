"""Check accepted sections against visible Markdown, independently of the renderer.

The source is authoritative. Frozen copies of legal paragraphs would prevent
accepted decisions from changing them. Candidate checks explicitly supply the
candidate source; ordinary CI uses the source from its own checkout.
"""
from __future__ import annotations

import os
import re
import sys
import unittest
from pathlib import Path

from docx import Document


DEFAULT_DOCUMENT = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)


def visible_line(line: str) -> str:
    line = re.sub(r"<!--.*?-->", "", line)
    line = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", line)
    return " ".join(line.replace("**", "").split())


def source_section(source: str, number: int) -> list[str]:
    matches = list(re.finditer(
        rf"^### § {number} — (.*?) <!-- section:[^>]+ -->$", source, re.MULTILINE,
    ))
    if len(matches) != 1:
        raise ValueError(f"Źródło musi zawierać dokładnie jeden § {number}")
    heading = matches[0]
    body = re.split(r"^#{2,3} ", source[heading.end():], maxsplit=1, flags=re.MULTILINE)[0]
    lines = [line for line in body.splitlines() if line.strip()]
    if not lines or any('<!-- unit:' not in line for line in lines):
        raise ValueError(f"Nieobsługiwana struktura treści § {number}; wymaga kontroli testu")
    return [f"§ {number}", visible_line(heading[1]), *map(visible_line, lines)]


def section(document: Document, start: str, end: str) -> list[str]:
    paragraphs = document.paragraphs
    start_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph.text == start)
    end_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph.text == end)
    return [paragraph.text for paragraph in paragraphs[start_index:end_index] if paragraph.text]


class ApprovedContentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path = (
            Path(sys.argv[1])
            if len(sys.argv) == 2 and sys.argv[1].lower().endswith(".docx")
            else DEFAULT_DOCUMENT
        )
        cls.document = Document(path)
        cls.source = Path(os.environ.get(
            'REGULAMIN_SOURCE_PATH', str(DEFAULT_DOCUMENT.with_suffix('.md')),
        )).read_text(encoding='utf-8')

    def test_approved_purpose_is_literal_and_unique(self):
        lines = [line for line in self.source.splitlines() if '<!-- unit:cel-glowny -->' in line]
        self.assertEqual(len(lines), 1)
        expected = visible_line(lines[0])
        matches = [p for p in self.document.paragraphs if p.text == expected]
        self.assertEqual(len(matches), 1, "Cel Regulaminu w DOCX musi odpowiadać źródłu")

    def test_section_2_defines_all_abbreviations_and_terms(self):
        self.assertEqual(section(self.document, "§ 2", "§ 3"), source_section(self.source, 2))

    def test_section_3_matches_approved_purpose_and_principles(self):
        self.assertEqual(section(self.document, "§ 3", "Rozdział 2"), source_section(self.source, 3))


if __name__ == "__main__":
    unittest.main(argv=[sys.argv[0]], verbosity=2)
