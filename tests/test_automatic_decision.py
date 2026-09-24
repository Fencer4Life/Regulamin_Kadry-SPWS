import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from narzedzia.create_decision_from_discussion import create, parse_discussion_form


BODY = '''### Problem
Problem źródłowy.
### Uzasadnienie
Uzgodnione uzasadnienie z dyskusji.
### Fragment Markdown do zastąpienia
   1) <!-- unit:sample --> Stara treść.
### Nowe brzmienie Markdown
   1) <!-- unit:sample --> Nowa treść.
### Zależy od
DR-002
'''


def discussion(body=BODY):
    return dict(html_url='https://github.com/Fencer4Life/Regulamin_Kadry-SPWS/discussions/99',
                title='Próba', created_at='2026-09-24T10:00:00Z', body=body,
                user={'login': 'Fencer4Life'})


class AutomaticDecisionTests(unittest.TestCase):
    def generate(self, root, body=BODY):
        event = root / 'event.json'
        event.write_text(json.dumps({'discussion': discussion(body)}))
        return create(event, root)

    def test_card_has_only_automatic_content_no_checkboxes_or_placeholders(self):
        with tempfile.TemporaryDirectory() as directory:
            card = self.generate(Path(directory)).read_text()
        for text in ('schema_version: 2', 'Uzgodnione uzasadnienie z dyskusji.',
                     '## Stary fragment Markdown', '## Nowy fragment Markdown', 'DR-002'):
            self.assertIn(text, card)
        for text in ('- [ ]', '- [x]', 'Do uzupełnienia', '## Problem', '## Kontekst',
                     '## Konsekwencje', '## Plan oceny', '## Odrzucone alternatywy'):
            self.assertNotIn(text, card)
        from narzedzia.decision_patch import read_fragments
        old, new = read_fragments(card)
        self.assertTrue(old.startswith('   1)'))
        self.assertIn('Nowa treść.', new)

    def test_missing_rationale_or_one_fragment_fails_without_creating_card(self):
        for body in (BODY.replace('Uzgodnione uzasadnienie z dyskusji.', '_No response_'),
                     BODY.replace('   1) <!-- unit:sample --> Nowa treść.', ''),
                     BODY.replace('Nowa treść.', 'Stara treść.')):
            with self.subTest(body=body), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                with self.assertRaises(ValueError):
                    self.generate(root, body)
                self.assertEqual(list(root.glob('DR-*.md')), [])

    def test_duplicate_field_is_rejected_not_guessed(self):
        with self.assertRaisesRegex(ValueError, 'Uzasadnienie'):
            parse_discussion_form(BODY + '\n### Uzasadnienie\nInny tekst\n')

    def test_unfinished_ambiguous_discussion_does_not_break_public_list(self):
        from narzedzia.sync_discussions import normalize_discussions
        node = {'number': 99, 'title': 'Próba', 'url': 'https://example/99',
                'createdAt': '2026-09-24T10:00:00Z', 'closedAt': None,
                'body': BODY + '\n### Uzasadnienie\nDrugi powód\n',
                'author': {'login': 'Fencer4Life'},
                'category': {'slug': 'propozycje-zmian-regulaminu'},
                'comments': {'totalCount': 0}, 'labels': {'nodes': []}}
        snapshot = normalize_discussions({'data': {'repository': {'discussions': {'nodes': [node]}}}}, 'now')
        self.assertEqual(snapshot['open_items'][0]['number'], 99)

    def test_decision_without_document_change_needs_explicit_resolution(self):
        body = '### Uzasadnienie\nWyjaśnienie.\n### Proponowane rozwiązanie\nZachowujemy obecny zapis.\n'
        with tempfile.TemporaryDirectory() as directory:
            card = self.generate(Path(directory), body).read_text()
        self.assertIn('Zachowujemy obecny zapis.', card)
        self.assertIn('zmiana_regulaminu: false', card)

    def test_pending_card_blocks_merge(self):
        from narzedzia.validate_regulation_change import validate_pending_decisions
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            cards = root / '_decyzje'
            cards.mkdir()
            card = self.generate(cards)
            changes = {str(card.relative_to(root)): 'A'}
            with self.assertRaisesRegex(ValueError, 'wdrażaj'):
                validate_pending_decisions(changes, root)

    def test_no_change_card_can_be_merged_without_docx(self):
        from narzedzia.validate_regulation_change import validate_pending_decisions
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            cards = root / '_decyzje'
            cards.mkdir()
            card = self.generate(cards, '### Uzasadnienie\nPowód.\n### Proponowane rozwiązanie\nBez zmian.\n')
            validate_pending_decisions({str(card.relative_to(root)): 'A'}, root)

    def test_refresh_uses_discussion_not_hand_edited_card(self):
        from narzedzia.pr_decision import refresh_from_discussion
        with tempfile.TemporaryDirectory() as directory:
            card = self.generate(Path(directory))
            current = discussion(BODY.replace('Nowa treść.', 'Aktualna treść.'))
            with patch('narzedzia.pr_decision.load_discussion', return_value=current):
                refresh_from_discussion(card, 'Fencer4Life/Regulamin_Kadry-SPWS')
            self.assertIn('Aktualna treść.', card.read_text())
            self.assertNotIn('Nowa treść.', card.read_text())

    def test_generated_card_has_no_trailing_whitespace_outside_code(self):
        with tempfile.TemporaryDirectory() as directory:
            card = self.generate(Path(directory), BODY.replace('dyskusji.', 'dyskusji.   ')).read_text()
        self.assertNotIn('dyskusji.   \n', card)

    def test_complete_discussion_to_docx_and_merge_gate_without_manual_card_edit(self):
        from narzedzia.pr_decision import apply_once
        from narzedzia.pr_docx_link import SOURCE, DOCX
        from narzedzia.validate_regulation_change import validate_pending_decisions
        from narzedzia.source_evidence import coverage
        from docx import Document
        repository = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, output = root / SOURCE, root / DOCX
            source.parent.mkdir()
            source.write_bytes((repository / SOURCE).read_bytes())
            old = next(line for line in source.read_text().splitlines()
                       if line.startswith('<!-- unit:publikacja-rankingu-zrodlo -->'))
            new = old.removesuffix('.') + ' [TEST AUTO].'
            body = ('### Uzasadnienie\nTest deterministycznej zmiany.\n'
                    '### Fragment Markdown do zastąpienia\n' + old + '\n'
                    '### Nowe brzmienie Markdown\n' + new + '\n')
            cards = root / '_decyzje'
            cards.mkdir()
            card = self.generate(cards, body)
            changes = {str(card.relative_to(root)): 'A', SOURCE: 'M', DOCX: 'M'}
            with self.assertRaisesRegex(ValueError, 'wdrażaj'):
                validate_pending_decisions(changes, root)
            self.assertTrue(apply_once(card, source, output))
            validate_pending_decisions(changes, root)
            self.assertFalse(apply_once(card, source, output))
            self.assertEqual(sum('[TEST AUTO]' in p.text for p in Document(output).paragraphs), 1)
            self.assertEqual(coverage(source, output)['missing'], [])

    def test_applied_card_is_never_refreshed_or_applied_twice(self):
        from narzedzia.pr_decision import refresh_from_discussion
        with tempfile.TemporaryDirectory() as directory:
            card = self.generate(Path(directory))
            card.write_text(card.read_text() + '\n<!-- applied-source-sha256:' + 'a' * 64 + ' -->\n')
            original = card.read_bytes()
            with patch('narzedzia.pr_decision.load_discussion') as load:
                refresh_from_discussion(card, 'Fencer4Life/Regulamin_Kadry-SPWS')
                load.assert_not_called()
            self.assertEqual(card.read_bytes(), original)

    def test_discussion_cannot_forge_applied_marker(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, 'zastrzeżony'):
                self.generate(Path(directory), BODY + '\n<!-- applied-source-sha256:' + 'a' * 64 + ' -->\n')

    def test_marker_without_changed_document_does_not_unlock_merge(self):
        from narzedzia.validate_regulation_change import validate_pending_decisions
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            cards = root / '_decyzje'
            cards.mkdir()
            card = self.generate(cards)
            card.write_text(card.read_text() + '\n<!-- applied-source-sha256:' + 'a' * 64 + ' -->\n')
            with self.assertRaisesRegex(ValueError, 'Markdown i DOCX'):
                validate_pending_decisions({str(card.relative_to(root)): 'M'}, root)

    def test_both_workflows_can_read_discussion_and_ci_does_not_skip_gate(self):
        root = Path(__file__).resolve().parents[1]
        for name in ('apply-decision-label.yml', 'apply-decision.yml'):
            text = (root / '.github/workflows' / name).read_text()
            self.assertIn('discussions: read', text)
        ci = (root / '.github/workflows/validate.yml').read_text()
        self.assertNotIn("if: github.event_name != 'workflow_dispatch'", ci)
        self.assertIn('python -m narzedzia.validate_regulation_change', ci)
