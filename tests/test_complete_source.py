import re
import unittest

from docx import Document
from tests.test_docx_current_contract import CANONICAL_SOURCE, CURRENT_DOCUMENT


def visible_source(text):
    text = text.split('\n+++\n', 1)[1]
    text = re.sub(r'<!--.*?-->', '', text, flags=re.S)
    text = re.sub(r'\[(?:unit|section|chapter|status):[^]]+\]', '', text)
    text = re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', text)
    text = '\n'.join(line.replace('|', ' ') if line.startswith('|') else line for line in text.splitlines())
    text = text.replace('**', '')
    return ' '.join(text.split())


class CompleteSourceTests(unittest.TestCase):
    def test_every_visible_paragraph_and_cell_has_visible_source(self):
        source = visible_source(CANONICAL_SOURCE.read_text())
        document = Document(CURRENT_DOCUMENT)
        paragraphs = list(document.paragraphs)
        for table in document.tables:
            paragraphs.extend(p for row in table.rows for cell in row.cells for p in cell.paragraphs)
        for section in document.sections:
            for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
                paragraphs.extend(part.paragraphs)
        missing = [p.text for p in paragraphs if p.text.strip() and ' '.join(p.text.split()) not in source]
        self.assertEqual(missing, [], 'Treść DOCX nieobecna w widocznym Markdown')

    def test_source_has_document_sections_not_only_metadata_or_preview(self):
        text = CANONICAL_SOURCE.read_text().split('\n+++\n', 1)[1]
        for value in ('# Regulamin powoływania', 'Spis treści', 'Konstrukcja regulaminu',
                      '| Wersja dokumentu |', '| Rozdział | Tytuł | Zakres |',
                      'ZAŁĄCZNIK NR 1', 'Historia wersji', '| Wersja | Data | Zakres zmian | Status |'):
            self.assertIn(value, text)
        self.assertNotRegex(text, r'\{\{(?:ref|term|days):')
