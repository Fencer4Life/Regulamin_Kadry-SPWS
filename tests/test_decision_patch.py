import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from narzedzia.decision_patch import format_fragments, replace_exact, read_fragments, apply_decision
from tests.test_docx_current_contract import CANONICAL_SOURCE, CURRENT_DOCUMENT


class DecisionPatchTests(unittest.TestCase):
    def test_decision_before_after_changes_annex_table_in_docx_and_preview(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, docx, card = root/'source.md', root/'source.docx', root/'DR-999.md'
            source.write_bytes(CANONICAL_SOURCE.read_bytes())
            old = next(line for line in source.read_text().splitlines() if line.startswith('| 4 | 54,3 |'))
            new = old.replace('54,3', '54,4', 1)
            card.write_text('---\nstatus: przyjęta\n---\n'+format_fragments(old, new))
            apply_decision(card, source, docx)
            self.assertIn(new, source.read_text())
            self.assertIn(new, source.with_suffix('.podglad.md').read_text())
            self.assertTrue(any(t.cell(1, 1).text == '54,4' for t in Document(docx).tables if len(t.columns) == 11))

    def test_template_has_distinct_empty_code_fields_with_instructions(self):
        template = (CANONICAL_SOURCE.parents[1] / "szablony/nowa-decyzja.md").read_text()
        self.assertIn("## Stary fragment Markdown", template)
        self.assertIn("## Nowy fragment Markdown", template)
        self.assertIn("kanonicznego", template)
        with self.assertRaises(ValueError):
            read_fragments(template)

    def test_fragment_roundtrip_preserves_indent_and_headings(self):
        old = "### [section:test] Test\n   1) [unit:a] Dawne."
        new = "### [section:test] Test\n   1) [unit:a] Nowe."
        self.assertEqual(read_fragments(format_fragments(old, new)), (old, new))
        self.assertEqual(replace_exact("prefix\n"+old+"\nsuffix", old, new),
                         "prefix\n"+new+"\nsuffix")

    def test_missing_repeated_and_unchanged_fragments_are_rejected(self):
        for source, old, new in [("aa aa", "aa", "bb"), ("aaa", "aa", "bb"), ("xx", "aa", "bb"),
                                 ("aa", "aa", "aa"), ("aa", "", "bb")]:
            with self.assertRaises(ValueError):
                replace_exact(source, old, new)

    def test_failed_generation_preserves_source_docx_and_preview(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, docx, card = root/"source.md", root/"source.docx", root/"DR-999.md"
            source.write_bytes(CANONICAL_SOURCE.read_bytes())
            docx.write_bytes(CURRENT_DOCUMENT.read_bytes())
            preview = source.with_suffix(".podglad.md")
            preview.write_text("previous preview")
            card.write_text("---\nstatus: przyjęta\n---\n" + format_fragments(
                "weteraniszermierki.pl", "www.weteraniszermierki.pl"))
            original = [p.read_bytes() for p in (source, docx, preview)]
            with patch("narzedzia.decision_patch.normalize_and_build", side_effect=ValueError("build")):
                with self.assertRaises(ValueError):
                    apply_decision(card, source, docx)
            self.assertEqual(original, [p.read_bytes() for p in (source, docx, preview)])

    def test_accepted_card_builds_all_outputs_and_rejects_second_application(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, docx, card = root/"source.md", root/"source.docx", root/"DR-999.md"
            source.write_bytes(CANONICAL_SOURCE.read_bytes())
            old = "weteraniszermierki.pl"
            card.write_text("---\nstatus: przyjęta\n---\n" + format_fragments(old, "www."+old))
            apply_decision(card, source, docx)
            self.assertTrue(docx.is_file())
            self.assertIn("www."+old, source.read_text())
            self.assertTrue(source.with_suffix(".podglad.md").is_file())
            with self.assertRaises(ValueError):
                apply_decision(card, source, docx)

    def test_workflow_updates_existing_pr_and_has_no_llm_dependency(self):
        workflow = (CANONICAL_SOURCE.parents[1] / ".github/workflows/apply-decision.yml").read_text()
        for expected in ("workflow_dispatch:", "decision_patch", "prepare_regulamin verify",
                         "unittest discover", "gh pr view", "--draft", 'git add -- "$card" "$source" "$docx" "$preview"'):
            self.assertIn(expected, workflow)
        for forbidden in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY", "curl", "--force", "git push origin main"):
            self.assertNotIn(forbidden, workflow)

    def test_unaccepted_decision_cannot_modify_source(self):
        with tempfile.TemporaryDirectory() as directory:
            card=Path(directory)/"card.md"
            card.write_text("---\nstatus: projekt\n---\n"+format_fragments("a","b"))
            with self.assertRaises(ValueError):
                apply_decision(card, CANONICAL_SOURCE, Path(directory)/"out.docx")

    def test_ambiguous_or_unclosed_fields_are_rejected(self):
        for text in (format_fragments("old", "new") + format_fragments("other", "new"),
                     format_fragments("old", "new").rsplit("```", 1)[0]):
            with self.assertRaises(ValueError):
                read_fragments(text)

    def test_bad_structure_keeps_original_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, docx, card = root/"source.md", root/"source.docx", root/"card.md"
            source.write_bytes(CANONICAL_SOURCE.read_bytes())
            docx.write_bytes(CURRENT_DOCUMENT.read_bytes())
            card.write_text("---\nstatus: przyjęta\n---\n" + format_fragments(
                "[unit:publikacja-rankingu-zrodlo]", "[unit:INVALID-ID]"))
            before = source.read_bytes(), docx.read_bytes(), card.read_bytes()
            with self.assertRaises(ValueError):
                apply_decision(card, source, docx)
            self.assertEqual(before, (source.read_bytes(), docx.read_bytes(), card.read_bytes()))
