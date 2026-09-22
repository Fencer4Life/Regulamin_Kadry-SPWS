"""The live document contract must follow accepted source edits, not old totals."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from narzedzia.docx_parity import document_content_contract
from narzedzia.prepare_regulamin import normalize_and_build
from tests import test_docx_current_contract as contract_tests


class DocxContractEditTests(unittest.TestCase):
    def assert_current_contract(self, source, docx):
        case = contract_tests.CurrentDocxContractTests()
        case.document = Document(docx)
        case.contract = document_content_contract(docx)
        with patch('tests.test_docx_current_contract.CANONICAL_SOURCE', source):
            case.test_current_document_shape_matches_source()
            case.test_current_document_matches_the_canonical_markdown_source()

    def test_source_can_add_or_remove_a_paragraph(self):
        original = contract_tests.CANONICAL_SOURCE.read_text()
        line = next(line for line in original.splitlines() if '<!-- unit:dms-spor -->' in line)
        for replacement in ('', line + '\n\n8. <!-- unit:e2e-added --> Dodatkowy akapit testowy.'):
            with self.subTest(replacement=replacement), tempfile.TemporaryDirectory() as directory:
                source, docx = Path(directory) / 'source.md', Path(directory) / 'out.docx'
                source.write_text(original.replace(line, replacement, 1))
                normalize_and_build(source, docx)
                self.assert_current_contract(source, docx)

    def test_missing_docx_paragraph_without_source_edit_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            source, docx = Path(directory) / 'source.md', Path(directory) / 'out.docx'
            source.write_bytes(contract_tests.CANONICAL_SOURCE.read_bytes())
            normalize_and_build(source, docx)
            document = Document(docx)
            paragraph = next(p for p in document.paragraphs if 'W przypadku braku porozumienia' in p.text)
            paragraph._p.getparent().remove(paragraph._p)
            document.save(docx)
            with self.assertRaises(AssertionError):
                self.assert_current_contract(source, docx)

    def test_changed_text_with_same_paragraph_count_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            source, docx = Path(directory) / 'source.md', Path(directory) / 'out.docx'
            source.write_bytes(contract_tests.CANONICAL_SOURCE.read_bytes())
            normalize_and_build(source, docx)
            document = Document(docx)
            paragraph = next(p for p in document.paragraphs if 'W przypadku braku porozumienia' in p.text)
            paragraph.runs[0].text += ' NIEZGODNY TEKST'
            document.save(docx)
            with self.assertRaises(AssertionError):
                self.assert_current_contract(source, docx)
