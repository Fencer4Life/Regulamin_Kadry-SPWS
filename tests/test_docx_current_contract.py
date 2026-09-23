from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from narzedzia.build_regulamin_docx import build_document
from narzedzia.docx_parity import document_content_contract


ROOT = Path(__file__).resolve().parents[1]
CURRENT_DOCUMENT = ROOT / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)
CANONICAL_SOURCE = ROOT / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.md"
)
EXPECTED_HEADINGS = [
    "Spis treści",
    "Konstrukcja Regulaminu",
    "Postanowienia ogólne",
    "§ 1", "Przedmiot regulaminu",
    "§ 2", "Definicje",
    "§ 3", "Cel i zasady wyłaniania reprezentacji",
    "Ranking indywidualny",
    "§ 4", "Rola rankingu indywidualnego",
    "§ 5", "Zawody uwzględniane w rankingu",
    "§ 6", "Zasady obliczania punktów",
    "§ 7", "Publikacja rankingu",
    "§ 8", "Wyniki w połączonych kategoriach wiekowych",
    "§ 9", "Zawodnicy uwzględniani w rankingu",
    "Powołania do startów indywidualnych",
    "§ 10", "Zasady powołań indywidualnych",
    "§ 11", "Rezygnacja i zastępstwo",
    "Dobór składu drużyny",
    "§ 12", "Pula kandydatów do drużyny",
    "§ 13", "Kategorie wiekowe w drużynie",
    "§ 14", "Drużynowe Mistrzostwa Świata",
    "§ 15", "Drużynowe Mistrzostwa Europy",
    "§ 16", "Powołanie uzupełniające",
    "Terminarz procesu powoływania",
    "§ 17", "Terminy procesu",
    "Ocena regulaminu i doskonalenie metody",
    "§ 18", "Posezonowa ocena działania regulaminu",
    "Postanowienia końcowe",
    "§ 19", "Wejście w życie",
    "Tabela punktacji Pucharu Polski Weteranów w szermierce",
    "Miejsca 1–10 · stawka 4–20 zawodników",
    "Miejsca 1–10 · stawka 21–37 zawodników",
    "Miejsca 1–10 · stawka 38–40 zawodników",
    "Miejsca 11–20 · stawka 11–27 zawodników",
    "Miejsca 11–20 · stawka 28–40 zawodników",
    "Miejsca 21–30 · stawka 21–37 zawodników",
    "Miejsca 21–30 · stawka 38–40 zawodników",
    "Miejsca 31–40 · stawka 31–40 zawodników",
    "Historia wersji",
]


class CurrentDocxContractTests(unittest.TestCase):
    def test_toc_has_no_prototype_entry_or_word_instruction(self):
        for paragraph in self.document.paragraphs:
            if paragraph.style.name.lower().startswith("toc"):
                self.assertNotIn("Informacja o prototypie", paragraph._p.xml)
            self.assertNotIn("Aktualizuj tabelę", paragraph.text)
        text = "\n".join(p.text for p in self.document.paragraphs)
        self.assertNotIn("Informacja o prototypie", text)
        self.assertNotIn("Przyjęcie treści podczas prac redakcyjnych", text)
        self.assertIn("Regulamin przyjmuje Zarząd SPWS w drodze uchwały", text)
        self.assertNotIn("właściwy organ PZSz", text)
        self.assertIn("Zarząd SPWS", self.document.core_properties.comments)

    def test_age_definitions_and_european_reserve_are_explicit(self):
        texts = [p.text for p in self.document.paragraphs]
        definitions = "\n".join(texts[texts.index("§ 2"):texts.index("§ 3")])
        for definition in ("V1 – od 40 do 49 lat", "V2 – od 50 do 59 lat", "V3 – od 60 do 69 lat", "V4 – 70 lat i więcej", "31 grudnia"):
            self.assertIn(definition, definitions)
        self.assertIn("3. Na Drużynowe Mistrzostwa Europy powołuje się także zawodnika rezerwowego, wybieranego wyłącznie spośród zawodników należących do puli kandydatów do drużyny.", texts)
        categories = "\n".join(texts[texts.index("§ 13"):texts.index("§ 14")])
        self.assertIn("kategorie V1 i V2", categories)
        self.assertIn("kategorie V3 i V4", categories)

    @classmethod
    def setUpClass(cls):
        cls.document = Document(CURRENT_DOCUMENT)
        cls.contract = document_content_contract(CURRENT_DOCUMENT)

    def test_current_document_shape_matches_source(self):
        # Editorial decisions can add/remove units and table rows. The current
        # Markdown defines the expected shape, not the migration-era totals.
        with tempfile.TemporaryDirectory() as directory:
            candidate = Path(directory) / "candidate.docx"
            build_document(CANONICAL_SOURCE, candidate)
            expected = Document(candidate)
            self.assertEqual(len(self.document.paragraphs), len(expected.paragraphs))
            self.assertEqual(len(self.document.tables), len(expected.tables))
            self.assertEqual(
                [len(table.rows) for table in self.document.tables],
                [len(table.rows) for table in expected.tables],
            )
            self.assertEqual(len(self.contract["blocks"]),
                             len(document_content_contract(candidate)["blocks"]))

    def test_current_document_heading_order_is_explicit(self):
        headings = [
            paragraph.text
            for paragraph in self.document.paragraphs
            if paragraph.style.name in {"Heading 1", "Heading 2", "Paragraf", "Tytuł paragrafu"}
        ]
        self.assertEqual(headings, EXPECTED_HEADINGS)

    def test_current_document_matches_the_canonical_markdown_source(self):
        with tempfile.TemporaryDirectory() as directory:
            candidate = Path(directory) / "candidate.docx"
            build_document(CANONICAL_SOURCE, candidate)
            self.assertEqual(
                self.contract,
                document_content_contract(candidate),
                "Śledzony DOCX nie odpowiada kanonicznemu źródłu Markdown",
            )


if __name__ == "__main__":
    unittest.main(verbosity=2)
