from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from narzedzia.build_regulamin_docx import build_document
from narzedzia.docx_parity import (
    document_content_contract,
    document_layout_contract,
    normalized_package_sha256,
)
from tests.test_docx_current_contract import CURRENT_DOCUMENT
from tests.test_ztp_model import METADATA


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "regulamin" / (
    "Regulamin-powolywania-reprezentacji-Polski-weteranow-w-szermierce_2026.md"
)


class GeneratedDocxTests(unittest.TestCase):
    def test_renders_arabic_chapter_separate_paragraph_title_and_draft_label(self):
        source_text = METADATA + '''
## [chapter:ogolne] Ogólne
<!-- scope: zakres -->
### [section:zasady] Zasady
1. [unit:przyjete] Treść przyjęta.
2. [unit:szkic-a] [status:source-draft] Pierwszy szkic.
3. [unit:szkic-b] [status:source-draft] Drugi szkic.
'''
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.md"
            target = Path(directory) / "target.docx"
            source.write_text(source_text, encoding="utf-8")
            build_document(source, target)
            document = Document(target)
            texts = [paragraph.text for paragraph in document.paragraphs]
            self.assertIn("Rozdział 1", texts)
            self.assertIn("§ 1", texts)
            self.assertIn("Zasady", texts)
            self.assertNotIn("§ 1. Zasady", texts)
            self.assertEqual(texts.count("BRUDNOPIS ZE ŹRÓDŁA — DO OPRACOWANIA"), 1)
            draft_paragraphs = [
                paragraph for paragraph in document.paragraphs
                if "szkic" in paragraph.text.lower()
            ]
            self.assertEqual(len(draft_paragraphs), 2)
            self.assertTrue(
                all(
                    run.font.color.rgb is not None
                    and str(run.font.color.rgb) == "595959"
                    for paragraph in draft_paragraphs
                    for run in paragraph.runs
                    if run.text
                )
            )

    def test_markdown_source_builds_a_content_identical_candidate(self):
        with tempfile.TemporaryDirectory() as directory:
            candidate = Path(directory) / "candidate.docx"
            build_document(SOURCE, candidate)
            self.assertTrue(candidate.is_file())
            self.assertEqual(
                document_content_contract(candidate),
                document_content_contract(CURRENT_DOCUMENT),
            )
            self.assertEqual(
                document_layout_contract(candidate),
                document_layout_contract(CURRENT_DOCUMENT),
            )

    def test_build_never_overwrites_the_source_document_implicitly(self):
        with tempfile.TemporaryDirectory() as directory:
            candidate = Path(directory) / "candidate.docx"
            before = CURRENT_DOCUMENT.read_bytes()
            build_document(SOURCE, candidate)
            self.assertEqual(CURRENT_DOCUMENT.read_bytes(), before)

    def test_two_builds_have_the_same_normalized_package(self):
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / "first.docx"
            second = Path(directory) / "second.docx"
            build_document(SOURCE, first)
            build_document(SOURCE, second)
            self.assertEqual(
                normalized_package_sha256(first),
                normalized_package_sha256(second),
            )


if __name__ == "__main__":
    unittest.main(verbosity=2)
