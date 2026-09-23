from __future__ import annotations

import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DEFAULT_DOCUMENT = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)
MAX_ATOMIC_TABLE_ROWS = 18
SECTION_TITLE_STYLE = "Tytuł paragrafu"
STATUS_LABEL_STYLE = "Etykieta brudnopisu"


class PaginationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path = (
            Path(sys.argv[1])
            if len(sys.argv) == 2 and sys.argv[1].lower().endswith(".docx")
            else DEFAULT_DOCUMENT
        )
        cls.document = Document(path)

    def test_each_chapter_starts_on_new_page(self):
        chapters = [p for p in self.document.paragraphs if p.text.startswith("Rozdział ")]
        self.assertGreater(len(chapters), 0, "Dokument nie zawiera nagłówków rozdziałów")
        for chapter in chapters:
            with self.subTest(chapter=chapter.text):
                self.assertIs(
                    chapter.paragraph_format.page_break_before,
                    True,
                    f"{chapter.text} nie ma wymuszonego początku nowej strony",
                )

    def test_sections_do_not_force_a_new_page(self):
        sections = [p for p in self.document.paragraphs if p.text.startswith("§ ")]
        self.assertGreater(len(sections), 0, "Dokument nie zawiera paragrafów")
        for section in sections:
            with self.subTest(section=section.text):
                self.assertIsNot(
                    section.paragraph_format.page_break_before,
                    True,
                    f"{section.text} nie powinien wymuszać początku nowej strony",
                )

    def test_each_section_opening_package_is_kept_together(self):
        paragraphs = self.document.paragraphs
        section_indexes = [
            index for index, paragraph in enumerate(paragraphs) if paragraph.text.startswith("§ ")
        ]
        self.assertGreater(len(section_indexes), 0, "Dokument nie zawiera paragrafów")

        for section_index in section_indexes:
            section = paragraphs[section_index]
            title = paragraphs[section_index + 1]
            cursor = section_index + 2
            labels = []
            while paragraphs[cursor].style.name == STATUS_LABEL_STYLE:
                labels.append(paragraphs[cursor])
                cursor += 1
            first_unit = paragraphs[cursor]

            with self.subTest(section=section.text):
                self.assertEqual(title.style.name, SECTION_TITLE_STYLE)
                self.assertEqual(first_unit.style.name, "Normal")
                self.assertIs(section.paragraph_format.keep_with_next, True)
                self.assertIs(title.paragraph_format.keep_with_next, True)
                for label in labels:
                    self.assertIs(label.paragraph_format.keep_with_next, True)
                self.assertIs(first_unit.paragraph_format.keep_together, True)

    def test_first_unit_children_remain_splittable(self):
        paragraphs = self.document.paragraphs
        definitions = next(
            index for index, paragraph in enumerate(paragraphs) if paragraph.text == "§ 2"
        )
        first_point = next(
            paragraph
            for paragraph in paragraphs[definitions + 1 :]
            if paragraph.text.startswith("1) SPWS")
        )
        self.assertIsNot(first_point.paragraph_format.keep_together, True)

    def test_chapters_have_no_redundant_manual_page_break_before_them(self):
        paragraphs = self.document.paragraphs
        chapters = [
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.startswith("Rozdział ")
        ]
        self.assertGreater(len(chapters), 1, "Dokument nie zawiera kolejnych rozdziałów")

        for chapter_index in chapters[1:]:
            chapter = paragraphs[chapter_index]
            preceding = paragraphs[chapter_index - 1]
            page_breaks = [
                element
                for element in preceding._p.iter(qn("w:br"))
                if element.get(qn("w:type")) == "page"
            ]
            with self.subTest(chapter=chapter.text):
                self.assertEqual(
                    page_breaks,
                    [],
                    f"Przed {chapter.text} znajduje się zbędny ręczny podział strony",
                )

    def test_each_table_is_small_enough_to_fit_on_one_page(self):
        for index, table in enumerate(self.document.tables, start=1):
            with self.subTest(table=index, first_cell=table.rows[0].cells[0].text):
                self.assertLessEqual(
                    len(table.rows),
                    MAX_ATOMIC_TABLE_ROWS,
                    f"Tabela {index} ma {len(table.rows)} wierszy i może przejść na następną stronę",
                )

    def test_each_table_row_is_indivisible(self):
        missing = []
        for table_index, table in enumerate(self.document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                properties = row._tr.get_or_add_trPr()
                if properties.find(qn("w:cantSplit")) is None:
                    missing.append((table_index, row_index))
        self.assertEqual(missing, [], f"Podzielne wiersze tabel: {missing[:20]}")

    def test_table_rows_are_kept_together(self):
        missing = []
        for table_index, table in enumerate(self.document.tables, start=1):
            for row_index, row in enumerate(table.rows[:-1], start=1):
                for cell_index, cell in enumerate(row.cells, start=1):
                    for paragraph_index, paragraph in enumerate(cell.paragraphs, start=1):
                        if paragraph.paragraph_format.keep_with_next is not True:
                            missing.append((table_index, row_index, cell_index, paragraph_index))
        self.assertEqual(missing, [], f"Brak keep-with-next w tabelach: {missing[:20]}")

    def test_each_table_is_followed_by_an_empty_paragraph(self):
        missing = []
        for table_index, table in enumerate(self.document.tables, start=1):
            following = table._tbl.getnext()
            text = "" if following is None else "".join(
                element.text or "" for element in following.iter(qn("w:t"))
            )
            if following is None or following.tag != qn("w:p") or text:
                missing.append(table_index)
        self.assertEqual(missing, [], f"Brak pustego akapitu po tabelach: {missing}")

    def test_annex_still_contains_every_score_from_4_to_40(self):
        annex_tables = [
            table
            for table in self.document.tables
            if table.rows[0].cells[0].text == "Liczba zawodników w stawce"
        ]
        cells = set()
        for table in annex_tables:
            places = [int(cell.text) for cell in table.rows[0].cells[1:]]
            for row in table.rows[1:]:
                participants = int(row.cells[0].text)
                for place, cell in zip(places, row.cells[1:]):
                    if cell.text != "—":
                        cells.add((participants, place))
        expected = {
            (participants, place)
            for participants in range(4, 41)
            for place in range(1, participants + 1)
        }
        self.assertEqual(cells, expected)


if __name__ == "__main__":
    # Ścieżka dokumentu jest argumentem testów, nie nazwą testu dla unittest.
    if len(sys.argv) > 1 and sys.argv[1].lower().endswith(".docx"):
        document_argument = sys.argv.pop(1)
        sys.argv.append(document_argument)
    unittest.main(argv=[sys.argv[0]], verbosity=2)
