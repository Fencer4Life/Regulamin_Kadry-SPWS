import unittest

from docx import Document
from docx.oxml.ns import qn
from tests.test_docx_current_contract import CURRENT_DOCUMENT


class DraftResourceLinksTests(unittest.TestCase):
    def test_annex_has_clickable_unofficial_full_table_link(self):
        document = Document(CURRENT_DOCUMENT)
        texts = [p.text for p in document.paragraphs]
        annex = document.paragraphs[texts.index("ZAŁĄCZNIK NR 1"):texts.index("Historia wersji")]
        links = [document.part.rels[node.get(qn("r:id"))].target_ref
                 for paragraph in annex for node in paragraph._p.xpath("w:hyperlink")]
        self.assertEqual(links, ["https://fencer4life.github.io/spws-automated-ranklist/tabela-punktacji.html"])
        self.assertTrue(any("wersja nieoficjalna" in p.text for p in annex))

    def test_cover_has_no_decorative_rule_above_title(self):
        document = Document(CURRENT_DOCUMENT)
        cover_text = "\n".join(p.text for p in document.paragraphs[:8])
        self.assertNotIn("━━━━", cover_text)

    def test_three_clickable_urls_are_placed_between_sections_7_and_8(self):
        document = Document(CURRENT_DOCUMENT)
        texts = [p.text for p in document.paragraphs]
        paragraphs = document.paragraphs[texts.index("§ 7"):texts.index("§ 8")]
        links = []
        for paragraph in paragraphs:
            for link in paragraph._p.xpath("w:hyperlink"):
                relationship = document.part.rels[link.get(qn("r:id"))]
                self.assertTrue(relationship.is_external)
                self.assertEqual("".join(link.xpath("w:r/w:t/text()")), relationship.target_ref)
                links.append(relationship.target_ref)
        self.assertEqual(links, [
            "https://fencer4life.github.io/spws-automated-ranklist/",
            "https://fencer4life.github.io/spws-automated-ranklist/tabela-punktacji.html",
            "https://fencer4life.github.io/spws-automated-ranklist/kalkulator-punktow.html",
        ])
        note = "\n".join(p.text for p in paragraphs)
        self.assertIn("Materiały pomocnicze — wersja nieoficjalna", note)
        self.assertIn("Regulamin nie został jeszcze uchwalony", note)
        self.assertIn("weteraniszermierki.pl", note)
        self.assertIn("Ranking indywidualny oraz kalkulator punktów są publikowane na stronie internetowej.", note)
