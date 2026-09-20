import tempfile
import unittest
from pathlib import Path

from docx import Document
from narzedzia.build_regulamin_docx import build_document
from narzedzia.normalize_regulamin_markdown import normalize_source
from narzedzia.prepare_regulamin import normalize_and_build
from tests.test_ztp_model import METADATA


SOURCE = METADATA.replace('\n+++\n', '\n[milestones]\nopening = 90\ncards = 75\nproposal = 60\nevent = 0\n+++\n') + '''
## [chapter:terminarz] Terminarz
<!-- scope: terminy -->
### [section:terminy] Terminy procesu
1. [unit:termin] Na {{days:cards}} dni przed zawodami ({{term:cards}}) udostępnia się karty.

{{table:timeline-process}}

| Termin | Zdarzenie | Zakres |
| --- | --- | --- |
| {{term:opening}} | Otwarcie | Zamknięcie danych |
| {{term:cards}} | Karty | Deklaracje do {{term:cards}} |
| {{term:proposal}} | Propozycja | Przekazanie PZSz |
| {{term:event}} | Zawody | Pierwszy dzień |

{{table:process}}

| Termin | Zakres |
| --- | --- |
| {{term:cards}} | Udostępnienie kart |
'''


class NominationTimelineTests(unittest.TestCase):
    def test_one_date_change_updates_prose_timeline_and_table_without_images(self):
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder) / 'source.md', Path(folder) / 'out.docx'
            source.write_text(SOURCE.replace('cards = 75', 'cards = 74'), encoding='utf-8')
            normalize_and_build(source, output)
            self.assertEqual(source.read_text(encoding='utf-8'), normalize_source(source))
            self.assertIn('{{term:cards}}', source.read_text(encoding='utf-8'))
            document = Document(output)
            self.assertIn('1. Na 74 dni przed zawodami (T−74) udostępnia się karty.', [p.text for p in document.paragraphs])
            timeline = next(t for t in document.tables if t.cell(0, 0).text == 'T−90')
            self.assertEqual([c.text for c in timeline.rows[0].cells], ['T−90', 'T−74', 'T−60', 'T'])
            self.assertEqual(timeline.cell(2, 1).text, 'Deklaracje do T−74')
            self.assertTrue(any(t.cell(0, 0).text == 'Termin' and t.cell(1, 0).text == 'T−74' for t in document.tables))
            self.assertEqual(len(document.inline_shapes), 0)
            self.assertNotIn('{{', document.element.xml)

    def test_invalid_reference_or_order_preserves_previous_docx(self):
        for invalid in (
            SOURCE.replace('{{term:cards}}', '{{term:missing}}'),
            SOURCE.replace('cards = 75', 'cards = 95'),
            SOURCE.replace('cards = 75', 'cards = -1'),
            SOURCE.replace('cards = 75', 'cards = true'),
            SOURCE.replace('cards = 75', 'cards = 90'),
            SOURCE.replace('| {{term:cards}} | Udostępnienie kart |', '| {{term:cards}} |'),
        ):
            with self.subTest(invalid=invalid), tempfile.TemporaryDirectory() as folder:
                source, output = Path(folder) / 'source.md', Path(folder) / 'out.docx'
                source.write_text(SOURCE, encoding='utf-8')
                build_document(source, output)
                before = output.read_bytes()
                source.write_text(invalid, encoding='utf-8')
                with self.assertRaises(ValueError):
                    normalize_and_build(source, output)
                self.assertEqual(before, output.read_bytes())

    def test_current_regulation_has_accepted_schedule_and_ms_exception(self):
        from tests.test_docx_current_contract import CURRENT_DOCUMENT
        document = Document(CURRENT_DOCUMENT)
        texts = [p.text for p in document.paragraphs]
        start, end = texts.index('§ 17'), texts.index('§ 18')
        schedule = '\n'.join(texts[start:end])
        self.assertIn('od T−90 do T−75 włącznie', schedule)
        self.assertIn('Najpóźniej na 60 dni', schedule)
        self.assertIn('w dniu poprzedzającym turniej drużynowy', schedule)
        self.assertIn('Nie wyznacza się zawodnika rezerwowego', schedule)
        self.assertIn('na Drużynowe Mistrzostwa Europy obejmuje skład pięcioosobowej drużyny wraz ze wskazaniem zawodnika rezerwowego', schedule)
        self.assertNotIn('MIEJSCE DO UZUPEŁNIENIA', schedule)
        timeline = next(t for t in document.tables if t.cell(0, 0).text == 'T−90')
        self.assertEqual([c.text for c in timeline.rows[0].cells], ['T−90', 'T−75', 'T−60', 'T'])
