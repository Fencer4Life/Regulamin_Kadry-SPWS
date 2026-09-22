import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch


class PrDecisionTests(unittest.TestCase):
    def test_real_decision_fixture_accepts_previously_edited_cover_date(self):
        from tests.test_docx_current_contract import CANONICAL_SOURCE
        content = CANONICAL_SOURCE.read_text()
        old_row = next(line for line in content.splitlines() if line.startswith('| Data projektu |'))
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'edited.md'
            source.write_text(content.replace(old_row, '| Data projektu | INNA-DATA-E2E |'))
            with patch('tests.test_docx_current_contract.CANONICAL_SOURCE', source):
                self.test_real_decision_changes_source_and_docx_only_once()

    def test_label_pr_may_change_only_decision_and_document_data(self):
        from narzedzia.pr_decision import check_changed_paths
        from narzedzia.pr_docx_link import SOURCE, DOCX
        card = '_decyzje/DR-026-test.md'
        check_changed_paths([{'filename': p} for p in (card, SOURCE, DOCX)], card)
        for extra in ('.github/workflows/validate.yml', 'narzedzia/pr_decision.py', 'tests/test_fake.py'):
            with self.assertRaises(ValueError):
                check_changed_paths([{'filename': card}, {'filename': extra}], card)

    def test_ci_is_dispatched_automatically_with_actions_permission(self):
        root = Path(__file__).resolve().parents[1]
        workflow = (root / '.github/workflows/apply-decision-label.yml').read_text()
        self.assertIn('actions: write', workflow)
        from narzedzia.pr_decision import dispatch_ci
        with patch('narzedzia.pr_decision.command') as command:
            dispatch_ci('owner/repo', 'decision/test')
            command.assert_called_once_with(['gh', 'workflow', 'run', 'validate.yml',
                                             '--repo', 'owner/repo', '--ref', 'decision/test'])

    def test_real_decision_changes_source_and_docx_only_once(self):
        from narzedzia.pr_decision import apply_once
        from narzedzia.decision_patch import format_fragments
        from tests.test_docx_current_contract import CANONICAL_SOURCE
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            card, source, docx = (root / p for p in ('card.md', 'source.md', 'out.docx'))
            source.write_bytes(CANONICAL_SOURCE.read_bytes())
            old_row = next(line for line in source.read_text().splitlines()
                           if line.startswith('| Data projektu |'))
            new_date = 'TEST-ETYKIETY' if old_row != '| Data projektu | TEST-ETYKIETY |' else 'TEST-ETYKIETY-2'
            card.write_text('---\nstatus: przyjęta\n---\n' + format_fragments(
                old_row, f'| Data projektu | {new_date} |'))
            self.assertTrue(apply_once(card, source, docx))
            before = [p.read_bytes() for p in (card, source, docx)]
            self.assertFalse(apply_once(card, source, docx))
            self.assertEqual(before, [p.read_bytes() for p in (card, source, docx)])
            self.assertEqual(Document(docx).tables[0].cell(2, 1).text, new_date)

    def test_selects_one_card_and_rejects_ambiguous_removed_or_nested_paths(self):
        from narzedzia.pr_decision import select_card
        card = {'filename': '_decyzje/DR-025-test.md', 'status': 'added'}
        self.assertEqual(select_card([card]), card['filename'])
        for files in ([], [card, {**card, 'filename': '_decyzje/DR-026-test.md'}],
                      [{**card, 'status': 'removed'}], [{**card, 'filename': '_decyzje/../DR-025-test.md'}]):
            with self.assertRaises(ValueError):
                select_card(files)

    def test_only_open_local_pr_with_label_can_apply(self):
        from narzedzia.pr_decision import check_pr
        pr = {'state': 'open', 'labels': [{'name': 'wdrażaj'}],
              'base': {'ref': 'main'}, 'head': {'sha': 'a' * 40, 'ref': 'decision/test', 'repo': {'full_name': 'owner/repo'}}}
        check_pr(pr, 'owner/repo')
        for invalid in ({**pr, 'state': 'closed'}, {**pr, 'labels': []},
                        {**pr, 'head': {**pr['head'], 'repo': {'full_name': 'other/fork'}}},
                        {**pr, 'head': {**pr['head'], 'ref': 'main'}}):
            with self.assertRaises(ValueError):
                check_pr(invalid, 'owner/repo')

    def test_repeated_label_verifies_but_never_reapplies_decision(self):
        from narzedzia.pr_decision import apply_once
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            card, source, docx = (root / p for p in ('card.md', 'source.md', 'out.docx'))
            card.write_text('<!-- applied-source-sha256:' + 'a' * 64 + ' -->')
            with patch('narzedzia.pr_decision.apply_decision') as apply, patch('narzedzia.pr_decision.verify') as verify:
                self.assertFalse(apply_once(card, source, docx))
                apply.assert_not_called()
                verify.assert_called_once_with(source, docx)

    def test_workflow_uses_trusted_code_not_pr_checkout(self):
        root = Path(__file__).resolve().parents[1]
        workflow = (root / '.github/workflows/apply-decision-label.yml').read_text()
        for required in ('pull_request_target:', 'types: [labeled]', "'wdrażaj'", 'ref: main', 'narzedzia.pr_decision'):
            self.assertIn(required, workflow)
        self.assertNotIn('ref: ${{ github.event.pull_request.head', workflow)
