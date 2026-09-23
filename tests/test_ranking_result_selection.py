from __future__ import annotations

import unittest
from pathlib import Path

from docx import Document


DOCUMENT = Path(__file__).resolve().parents[1] / "regulamin" / (
    "Regulamin-powolywania-Reprezentacji-Polski-Weteranow-w-szermierce_2026.docx"
)

EXPECTED_SELECTION_RULES = [
    "3. Łączna liczba punktów zawodnika w rankingu indywidualnym jest sumą punktów "
    "przypisanych do ośmiu pozycji wynikowych:",
    "1) trzech obowiązkowych pozycji krajowych;",
    "2) pięciu pozycji otwartych.",
    "4. Trzy obowiązkowe pozycje krajowe wypełnia się wynikami uzyskanymi w zawodach "
    "Pucharu Polski Weteranów w Szermierce oraz Mistrzostwach Polski Weteranów w Szermierce, "
    "według następujących zasad:",
    "1) spośród wszystkich wyników uzyskanych przez zawodnika w tych zawodach wybiera się "
    "trzy wyniki, za które zawodnik otrzymał najwyższą liczbę punktów rankingowych;",
    "2) jeżeli zawodnik uzyskał mniej niż trzy wyniki w Pucharze Polski Weteranów w Szermierce "
    "lub Mistrzostwach Polski Weteranów w Szermierce, za każdą niewypełnioną obowiązkową "
    "pozycję krajową przyjmuje się zero punktów rankingowych.",
    "5. Na pięć pozycji otwartych składa się pięć najwyżej punktowanych spośród pozostałych "
    "wyników zawodnika uwzględnianych w rankingu.",
    "6. Najpierw wybiera się wyniki wypełniające trzy obowiązkowe pozycje krajowe, a następnie "
    "wyniki wypełniające pięć pozycji otwartych.",
    "7. Żaden wynik nie może zostać policzony dwukrotnie.",
]


class RankingResultSelectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.paragraphs = [paragraph.text for paragraph in Document(DOCUMENT).paragraphs]
        start = cls.paragraphs.index("§ 5")
        end = cls.paragraphs.index("§ 6")
        cls.section = [text for text in cls.paragraphs[start:end] if text]

    def test_section_contains_the_complete_eight_result_selection_rule(self):
        self.assertEqual(self.section[-len(EXPECTED_SELECTION_RULES) :], EXPECTED_SELECTION_RULES)

    def test_redundant_published_data_requirement_is_removed(self):
        self.assertNotIn(
            "8) opublikowane wyniki zawierają dane niezbędne do ich zweryfikowania oraz obliczenia "
            "punktów rankingowych.",
            self.section,
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
