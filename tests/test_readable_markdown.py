import tempfile
import unittest
from pathlib import Path

from narzedzia.docx_model import parse_regulation_source
from narzedzia.normalize_regulamin_markdown import normalize_source
from tests.test_docx_current_contract import CANONICAL_SOURCE, CURRENT_DOCUMENT


class ReadableMarkdownTests(unittest.TestCase):
    def test_annex_tables_and_cover_values_are_authoritative_source(self):
        from docx import Document
        from narzedzia.build_regulamin_docx import build_document
        model = parse_regulation_source(CANONICAL_SOURCE)
        self.assertTrue(model.annex_tables)
        reference = Document(CURRENT_DOCUMENT)
        self.assertEqual(model.metadata['cover_version'], reference.tables[0].cell(0, 1).text)
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory)/'source.md', Path(directory)/'out.docx'
            text = CANONICAL_SOURCE.read_text()
            text = text.replace('| Wersja dokumentu | 0.1 |', '| Wersja dokumentu | TEST-VERSION |')
            text = text.replace('| 4 | 54,3 |', '| 4 | 9,9 |', 1)
            source.write_text(text)
            build_document(source, output)
            document = Document(output)
            self.assertEqual(document.tables[0].cell(0, 1).text, 'TEST-VERSION')
            self.assertTrue(any(t.cell(1, 1).text == '9,9' for t in document.tables if len(t.columns) == 11))

    def test_links_live_in_source_body_and_annex_not_metadata(self):
        text = CANONICAL_SOURCE.read_text()
        metadata, body = text[4:].split("\n+++\n", 1)
        self.assertNotIn("resource_", metadata)
        section = body.split("<!-- section:publikacja-rankingu -->", 1)[1].split("### ", 1)[0]
        self.assertIn("<!-- note:", section)
        self.assertIn("[https://fencer4life.github.io/spws-automated-ranklist/]", section)
        self.assertIn("tabela-punktacji.html", body.split("{{annex:points}}", 1)[1])
        self.assertEqual(normalize_source(CANONICAL_SOURCE), text)

    def test_notes_have_roundtrip_and_no_invented_numbering(self):
        model = parse_regulation_source(CANONICAL_SOURCE)
        self.assertEqual(len(model.annex_notes), 1)
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.md"
            source.write_text(normalize_source(CANONICAL_SOURCE))
            self.assertEqual(normalize_source(source), CANONICAL_SOURCE.read_text())

    def test_preview_contains_every_table_and_clickable_links(self):
        from docx import Document
        from narzedzia.markdown_preview import render_preview
        preview = render_preview(CURRENT_DOCUMENT)
        for table in Document(CURRENT_DOCUMENT).tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        self.assertIn(cell.text.replace("\n", "<br>"), preview)
        self.assertIn("](https://fencer4life.github.io/spws-automated-ranklist/)", preview)
        self.assertNotIn("{{annex:", preview)
